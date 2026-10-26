"""카드 장면(scene) 목록 → 편집 가능한 PPTX.

모든 글자는 텍스트 상자, 도형·선·배경은 PowerPoint 도형, 비교표는 PowerPoint 표, 막대 그래프는 임베디드 엑셀 데이터가 있는
PowerPoint 차트(데이터 편집 가능)로 만든다. 슬라이드 전체를 이미지 한 장으로 넣지 않는다.
"""

import io

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Pt

from app.services.card_news_scene import PAD_CM, _lum

STATE_TEXT = {"verified": "연구원 검증됨", "not_verified": "근거 미검증", "stale": "검증 후 수정됨·재검증 필요"}
ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
VALIGN = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.upper())


def _set_fonts(run, font: str, size: float, bold: bool, color: str) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):  # 한글(동아시아) 글꼴도 지정해야 PowerPoint에서 한글이 같은 글꼴로 표시된다
        for old in rpr.findall(qn(tag)):
            rpr.remove(old)
        el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", font)


def _fill_line(shape, fill, line, line_w=0.75) -> None:
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(fill)
    else:
        shape.fill.background()
    if line:
        shape.line.color.rgb = _rgb(line)
        shape.line.width = Pt(line_w)
    else:
        shape.line.fill.background()
    shape.shadow.inherit = False
    style = shape._element.find(qn("p:style"))
    if style is not None:  # 도형 스타일의 그림자 참조까지 끈다(평면 도형)
        eff = style.find(qn("a:effectRef"))
        if eff is not None:
            eff.set("idx", "0")


def _add_text(slide, it: dict) -> None:
    tb = slide.shapes.add_textbox(Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]))
    tb.name = it.get("key") or it.get("name") or "label"
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = Cm(PAD_CM)
    tf.margin_top = tf.margin_bottom = Cm(0.05)
    tf.vertical_anchor = VALIGN[it["valign"]]
    if it.get("links") and _lum(it["color"]) > 0.45:
        # 어두운 배경 카드: 하이퍼링크 글자는 PowerPoint 테마 링크색(파랑)으로 바뀌어 읽기 어려우므로 글자색은 유지하고
        # 텍스트 상자 전체에 첫 출처 링크를 건다(전체 출처 URL은 슬라이드 노트에 있음)
        tb.click_action.hyperlink.address = it["links"][0]["url"]
        it = dict(it, links=None)
    if it.get("links"):  # 출처 문구: 원문으로 이동하는 하이퍼링크(PNG는 클릭할 수 없으므로 출처 목록은 슬라이드 노트·별도 파일에 둔다)
        p = tf.paragraphs[0]
        p.alignment = ALIGN[it["align"]]
        p.line_spacing = it.get("spacing", 1.2)
        run = p.add_run()
        run.text = "출처: "
        _set_fonts(run, it["font"], it["size"], it["bold"], it["color"])
        for n, lk in enumerate(it["links"]):
            if n:
                sep = p.add_run()
                sep.text = ", "
                _set_fonts(sep, it["font"], it["size"], it["bold"], it["color"])
            r = p.add_run()
            r.text = lk["label"]
            _set_fonts(r, it["font"], it["size"], it["bold"], it["color"])
            r.hyperlink.address = lk["url"]
        rest = str(it["text"])[len("출처: " + ", ".join(lk["label"] for lk in it["links"])):]
        if rest:  # "외 n건" 등 링크가 아닌 꼬리 문구
            tail = p.add_run()
            tail.text = rest
            _set_fonts(tail, it["font"], it["size"], it["bold"], it["color"])
        return
    for idx, para_text in enumerate(str(it["text"]).split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = ALIGN[it["align"]]
        p.line_spacing = it.get("spacing", 1.2)
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        run = p.add_run()
        run.text = para_text
        _set_fonts(run, it["font"], it["size"], it["bold"], it["color"])


def _cell_border(cell, color: str, width_pt: float = 0.75) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for old in tcpr.findall(qn(tag)):
            tcpr.remove(old)
    for pos, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
        ln = etree.Element(qn(tag), w=str(int(width_pt * 12700)), cap="flat", cmpd="sng", algn="ctr")
        solid = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(solid, qn("a:srgbClr"), val=color.upper())
        etree.SubElement(ln, qn("a:prstDash"), val="solid")
        tcpr.insert(pos, ln)  # 테두리는 채우기 앞에 와야 스키마에 맞는다


def _add_table(slide, it: dict) -> None:
    rows = it["rows"]
    n_cols = len(it["col_widths"])
    gf = slide.shapes.add_table(len(rows), n_cols, Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]))
    gf.name = it.get("key") or "table"
    table = gf.table
    table.first_row = False
    table.horz_banding = False
    for c, w in enumerate(it["col_widths"]):
        table.columns[c].width = Cm(w)
    for r, row in enumerate(rows):
        table.rows[r].height = Cm(row["h"])
        for c, cell_def in enumerate(row["cells"]):
            cell = table.cell(r, c)
            cell.margin_left = cell.margin_right = Cm(0.12)
            cell.margin_top = cell.margin_bottom = Cm(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(cell_def["fill"])
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = ALIGN[cell_def.get("align", "left")]
            p.line_spacing = 1.15
            run = p.add_run()
            run.text = cell_def["text"]
            _set_fonts(run, it["font"], cell_def["size"], cell_def.get("bold", False), cell_def["color"])
            _cell_border(cell, it["border"])


def _add_chart(slide, it: dict) -> None:
    data = CategoryChartData()
    data.categories = it["categories"]
    data.add_series(f"{it['title'] or '값'} ({it['unit']})", it["values"])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]), data)
    gf.name = "chart"
    chart = gf.chart
    chart.has_legend = False
    chart.font.size = Pt(8)
    chart.font.name = it["font"]
    chart.has_title = bool(it["title"])
    if it["title"]:
        chart.chart_title.text_frame.text = it["title"]
        run = chart.chart_title.text_frame.paragraphs[0].runs[0]
        run.font.size = Pt(8.5)
        run.font.bold = True
        run.font.color.rgb = _rgb(it["navy"])
    plot = chart.plots[0]
    plot.gap_width = 70
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size = Pt(8)
    labels.font.bold = True
    labels.number_format = "#,##0.##"
    labels.number_format_is_linked = False
    labels.position = XL_LABEL_POSITION.OUTSIDE_END
    series = plot.series[0]
    series.format.fill.solid()
    series.format.fill.fore_color.rgb = _rgb(it["color"])
    chart.value_axis.has_major_gridlines = False
    chart.value_axis.visible = False
    chart.category_axis.tick_labels.font.size = Pt(8)
    chart.category_axis.format.line.color.rgb = _rgb("9AA3B2")


def build_pptx(scenes: list[dict], title: str, font: str) -> bytes:
    """scenes: 카드별 장면(for_export=True로 만든 것). 모든 카드는 같은 크기(첫 장면 기준)."""
    prs = Presentation()
    prs.slide_width = Cm(scenes[0]["width"])
    prs.slide_height = Cm(scenes[0]["height"])
    prs.core_properties.title = title
    prs.core_properties.author = "TNP AI Platform"
    prs.core_properties.comments = "카드뉴스 — 텍스트·도형·표·차트는 PowerPoint에서 직접 편집 가능한 객체입니다(전체 이미지 삽입 아님)."
    blank = prs.slide_layouts[6]
    for scene in scenes:
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(scene["background"])
        if scene.get("sources"):
            notes = ["출처 목록(근거 상태는 내부 검증 기록 기준이며 IITP 공식 확인이 아닙니다)"]
            for s in scene["sources"]:
                notes.append(" · ".join(x for x in (s["label"], s.get("doc_title"), f"발표일 {s['published']}" if s.get("published") else None,
                                                    s["url"], STATE_TEXT.get(s["state"], s["state"])) if x))
            slide.notes_slide.notes_text_frame.text = "\n".join(notes)
        for it in scene["items"]:
            kind = it["type"]
            if kind == "rect":
                radius = it.get("radius") or 0
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]))
                if radius:
                    shape.adjustments[0] = min(0.5, radius / max(min(it["w"], it["h"]), 0.01))
                _fill_line(shape, it.get("fill"), it.get("line"))
                shape.name = it.get("name") or "box"
            elif kind == "ellipse":
                shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(it["x"]), Cm(it["y"]), Cm(it["w"]), Cm(it["h"]))
                _fill_line(shape, it.get("fill"), None)
                shape.name = it.get("name") or "dot"
            elif kind == "line":
                conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(it["x1"]), Cm(it["y1"]), Cm(it["x2"]), Cm(it["y2"]))
                conn.line.color.rgb = _rgb(it["color"])
                conn.line.width = Pt(it["width_pt"])
                conn.name = "line"
                style = conn._element.find(qn("p:style"))
                if style is not None:  # 선에도 그림자 참조를 끈다
                    eff = style.find(qn("a:effectRef"))
                    if eff is not None:
                        eff.set("idx", "0")
            elif kind == "text":
                _add_text(slide, it)
            elif kind == "table":
                _add_table(slide, it)
            elif kind == "chart":
                _add_chart(slide, it)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
