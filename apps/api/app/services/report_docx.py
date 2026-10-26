"""보고서 스냅샷(JSON) → 실제 편집 가능한 DOCX 생성 엔진.

- python-docx에는 각주 API가 없어서 word/footnotes.xml 파트를 직접 만들어 연결한다.
  → 출력 파일의 각주는 위첨자 흉내가 아니라 Word의 '실제 각주'(Footnote)다.
- 목차는 Word TOC 필드다. 페이지 번호는 PDF를 실제로 렌더링해서 알아낸 값(toc_pages)이 있을 때만
  캐시 결과로 넣고, 모르면 번호를 넣지 않는다(임의의 번호를 만들지 않는다).
- 표/각주/하이퍼링크/글머리표·번호목록/굵게·기울임·밑줄을 TipTap(ProseMirror) JSON에서 변환한다.
"""

import io
import re
from dataclasses import dataclass, field
from datetime import datetime

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Mm, Pt, RGBColor

ROMAN = ["Ⅰ", "Ⅱ", "Ⅲ", "Ⅳ", "Ⅴ", "Ⅵ", "Ⅶ", "Ⅷ", "Ⅸ", "Ⅹ"]

DEFAULT_TEMPLATE_CONFIG = {
    "font_body": "맑은 고딕",
    "font_heading": "맑은 고딕",
    "font_size_body": 10.5,
    "font_size_footnote": 8.5,
    "line_spacing": 1.4,
    "para_space_after_pt": 4,
    "margins_mm": {"top": 25, "bottom": 22, "left": 22, "right": 22},
    "header_text": "주간 AI·ICT 동향",
    "show_page_numbers": True,
    "color_primary": "1F3A5F",
    "color_table_header": "DCE4F0",
    "table_caption_format": "[표 {part}-{n}]",
    "footnote_format": "{org}, 「{title}」, {published}, {url} (열람 {accessed})",
    "footnote_short_format": "{org}, 앞의 각주 {n}번 참조",
    "footnote_show_verification_tag": True,
    "publisher": "테크노베이션파트너스(TNP)",
}


def merge_template_config(config: dict | None) -> dict:
    merged = dict(DEFAULT_TEMPLATE_CONFIG)
    merged["margins_mm"] = dict(DEFAULT_TEMPLATE_CONFIG["margins_mm"])
    if config:
        for key, value in config.items():
            if key == "margins_mm" and isinstance(value, dict):
                merged["margins_mm"].update(value)
            else:
                merged[key] = value
    return merged


@dataclass
class BuildResult:
    docx_bytes: bytes
    footnote_count: int = 0
    table_count: int = 0
    # (Word 제목 레벨, 제목 텍스트) — PDF에서 실제 페이지 번호를 찾는 데 사용
    headings: list[tuple[int, str]] = field(default_factory=list)
    part_headings: list[str] = field(default_factory=list)


def _set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size:
        run.font.size = Pt(size)


def _style_font(style, name: str, size: float | None = None, bold: bool | None = None, color: str | None = None) -> None:
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color:
        style.font.color.rgb = RGBColor.from_string(color)


def _title_key(text: str) -> str:
    """제목 비교용 정규화 — '[시범]' 같은 앞 머리표와 공백 차이를 무시한다."""
    return re.sub(r"\s+", "", re.sub(r"^(\s*\[[^\]]*\]\s*)+", "", text or ""))


def _nbsp(text: str) -> str:
    """영문 약어·기술명(예: FMVSS 305a)이 줄바꿈에서 중간에 끊기지 않도록 영숫자 사이 공백을 줄바꿈 금지 공백으로 바꾼다."""
    return re.sub(r"(?<=[A-Za-z0-9\)])\s(?=[A-Za-z0-9\(])", "\u00a0", text)


def _fmt_date(value: str | None) -> str:
    if not value:
        return "발표일 미확인"
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")).strftime("%Y.%m.%d.")
    except ValueError:
        return value[:10]


def format_footnote_text(fn: dict, config: dict, short_ref: int | None = None) -> str:
    """관리자가 template.config.footnote_format으로 바꿀 수 있는 출처 표기 규칙."""
    values = {
        "org": fn.get("source_org") or "",
        "title": fn.get("doc_title") or fn.get("anchor_text") or "",
        "published": _fmt_date(fn.get("published_at")) if fn.get("published_at") else "발표일 미확인",
        "accessed": _fmt_date(fn.get("accessed_at")) if fn.get("accessed_at") else "",
        "url": fn.get("source_url") or "",
        "location": fn.get("location_label") or "위치 미확인",
        "n": str(short_ref) if short_ref else "",
    }

    class _Safe(dict):
        def __missing__(self, key):
            return ""

    template = config["footnote_short_format"] if short_ref else config["footnote_format"]
    text = template.format_map(_Safe(values))
    text = re.sub(r"「\s*」", "", text)
    text = re.sub(r"\(\s*열람\s*\)", "", text)
    text = re.sub(r"(\s*,\s*){2,}", ", ", text)
    text = re.sub(r"^\s*,\s*", "", text).strip()
    if config.get("footnote_show_verification_tag", True):
        if fn.get("verification_status") != "matched":
            text += " [근거 미검증]"
        if not fn.get("location_label"):
            text += " [원문 내 위치 미확인]"
    return text


class DocxBuilder:
    def __init__(self, snapshot: dict, toc_pages: list[int | None] | None = None):
        self.snapshot = snapshot
        self.config = merge_template_config(snapshot.get("template", {}).get("config"))
        self.toc_pages = toc_pages
        self.doc = Document()
        self.footnotes: list[tuple[int, str, str | None]] = []
        self.compact = snapshot["report"].get("layout") == "compact"
        self._seen_sources: dict[str, int] = {}
        self.headings: list[tuple[int, str]] = []
        self.part_headings: list[str] = []
        self.table_count = 0
        self._fn_lookup: dict[str, dict] = {}
        self._table_counter_in_part = 0
        self._current_part_no = 0

    # ---- 문서 기본 설정 ----
    def _setup_styles(self) -> None:
        cfg = self.config
        styles = self.doc.styles
        body_font, head_font = cfg["font_body"], cfg["font_heading"]
        normal = styles["Normal"]
        _style_font(normal, body_font, cfg["font_size_body"])
        normal.paragraph_format.line_spacing = cfg["line_spacing"]
        normal.paragraph_format.space_after = Pt(cfg["para_space_after_pt"])
        sizes = {1: 17, 2: 13.5, 3: 11.5, 4: 10.5}
        for level, size in sizes.items():
            st = styles[f"Heading {level}"]
            _style_font(st, head_font, size, True, cfg["color_primary"])
            st.font.italic = False
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(14 if level == 1 else 10)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.line_spacing = 1.2
        for name in ("List Bullet", "List Bullet 2"):
            _style_font(styles[name], body_font, cfg["font_size_body"])
        if "Caption" in [s.name for s in styles]:
            _style_font(styles["Caption"], head_font, 9, True, cfg["color_primary"])
        title = styles["Title"]
        _style_font(title, head_font, 26, True, cfg["color_primary"])
        tppr = title.element.get_or_add_pPr()
        for border in tppr.findall(qn("w:pBdr")):  # 기본 Title 스타일의 파란 밑줄 제거(표지에 직접 규칙선을 둔다)
            tppr.remove(border)

        footnote_text = styles.add_style("Footnote Text", WD_STYLE_TYPE.PARAGRAPH)
        footnote_text.base_style = normal
        _style_font(footnote_text, body_font, cfg["font_size_footnote"])
        footnote_text.paragraph_format.line_spacing = 1.15
        footnote_text.paragraph_format.space_after = Pt(1)
        footnote_ref = styles.add_style("Footnote Reference", WD_STYLE_TYPE.CHARACTER)
        footnote_ref.font.superscript = True
        hyper = styles.add_style("Hyperlink", WD_STYLE_TYPE.CHARACTER)
        hyper.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        hyper.font.underline = True

        section = self.doc.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        m = cfg["margins_mm"]
        section.top_margin, section.bottom_margin = Mm(m["top"]), Mm(m["bottom"])
        section.left_margin, section.right_margin = Mm(m["left"]), Mm(m["right"])
        section.different_first_page_header_footer = not self.compact  # 표지에는 머리글·바닥글을 넣지 않는다(간결형은 첫 쪽에도 표시)

        # 머리글
        header_p = section.header.paragraphs[0]
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = header_p.add_run(cfg["header_text"] or "")
        _set_run_font(hr, body_font, 8.5)
        hr.font.color.rgb = RGBColor(0x66, 0x72, 0x85)
        # 바닥글: 상태 표기 + 페이지 번호(필드)
        footer_p = section.footer.paragraphs[0]
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = self._status_label()
        if label:
            lr = footer_p.add_run(label + "   ")
            _set_run_font(lr, body_font, 8)
            lr.font.color.rgb = RGBColor(0x9A, 0x4B, 0x00)
        if cfg["show_page_numbers"]:
            self._add_field(footer_p, "PAGE", "1", body_font, 9)

    def _status_label(self) -> str:
        rep = self.snapshot["report"]
        parts = []
        if rep.get("status") == "final":
            label = "출력용 확정본(예외 승인)" if rep.get("exception") else "출력용 확정본(내부 절차)"
            parts.append(label + ("" if rep.get("iitp_confirmed") else " · IITP 공식 확인 전"))
        elif not rep.get("iitp_confirmed"):
            parts.append("내부 검토용 초안 · IITP 공식 확인 전")
        if rep.get("is_pilot"):
            parts.append("시범 출력(개발용 테스트 자료)")
        return " / ".join(parts)

    def _add_field(self, paragraph, instruction: str, cached: str, font: str, size: float) -> None:
        def run_with(child):
            r = paragraph.add_run()
            _set_run_font(r, font, size)
            r._r.append(child)
            return r

        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        run_with(begin)
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instruction} "
        run_with(instr)
        sep = OxmlElement("w:fldChar")
        sep.set(qn("w:fldCharType"), "separate")
        run_with(sep)
        r = paragraph.add_run(cached)
        _set_run_font(r, font, size)
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        run_with(end)

    # ---- 표지·목차 ----
    def _title_size(self, text: str) -> float:
        n = len(text)
        return 26 if n <= 18 else 22 if n <= 30 else 19 if n <= 48 else 16

    def _build_cover(self) -> None:
        rep = self.snapshot["report"]
        cfg = self.config
        for _ in range(5):
            self.doc.add_paragraph()
        label = self.doc.add_paragraph()
        r = label.add_run(rep.get("kind_label", "주간 AI·ICT 동향"))
        _set_run_font(r, cfg["font_heading"], 12)
        r.font.color.rgb = RGBColor.from_string(cfg["color_primary"])
        r.bold = True
        title = self.doc.add_paragraph(style="Title")
        title.paragraph_format.space_after = Pt(6)
        tr = title.add_run(_nbsp(rep["title"]))
        _set_run_font(tr, cfg["font_heading"], self._title_size(rep["title"]))  # 긴 제목은 글자 크기를 줄여 단어가 어색하게 끊기지 않게 한다
        tr.bold = True
        tr.font.color.rgb = RGBColor.from_string(cfg["color_primary"])
        if rep.get("title_en"):
            en = self.doc.add_paragraph()
            er = en.add_run(_nbsp(rep["title_en"]))
            _set_run_font(er, cfg["font_heading"], 13)
            er.italic = True
            er.font.color.rgb = RGBColor(0x44, 0x4F, 0x62)
        rule = self.doc.add_paragraph()
        ppr = rule._p.get_or_add_pPr()
        pbdr = parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="{cfg["color_primary"]}"/></w:pBdr>'
        )
        ppr.append(pbdr)
        meta_lines = []
        if rep.get("edition_label"):
            meta_lines.append(f"발간 회차: {rep['edition_label']}")
        if rep.get("publish_date"):
            meta_lines.append(f"발간(예정)일: {_fmt_date(rep['publish_date'])}")
        if cfg.get("publisher"):
            meta_lines.append(f"작성: {cfg['publisher']}")
        meta_lines.append("문서 상태: " + ("출력용 확정본(내부 절차)" if rep.get("status") == "final" else "내부 검토용 초안")
                          + ("" if rep.get("iitp_confirmed") else " · IITP 공식 확인 전"))
        meta_lines.append(f"출력 기준 템플릿: {self.snapshot['template']['name']} v{self.snapshot['template']['version']}"
                          + (" (임시 표준 템플릿 — IITP 공식 양식 아님)" if self.snapshot["template"].get("is_provisional") else ""))
        for line in meta_lines:
            p = self.doc.add_paragraph()
            rr = p.add_run(line)
            _set_run_font(rr, cfg["font_body"], 11)
        for _ in range(2):
            self.doc.add_paragraph()
        notice = self._status_label()
        if notice:
            p = self.doc.add_paragraph()
            rr = p.add_run("※ " + notice)
            _set_run_font(rr, cfg["font_body"], 10)
            rr.font.color.rgb = RGBColor(0x9A, 0x4B, 0x00)
            rr.bold = True
        self._page_break()

    def _build_compact_title(self) -> None:
        """간결형: 표지·목차 페이지 없이 첫 쪽 상단에 제목 블록만 둔다."""
        rep = self.snapshot["report"]
        cfg = self.config
        title = self.doc.add_paragraph(style="Title")
        title.paragraph_format.space_after = Pt(2)
        tr = title.add_run(_nbsp(rep["title"]))
        _set_run_font(tr, cfg["font_heading"], min(self._title_size(rep["title"]), 20))
        tr.bold = True
        tr.font.color.rgb = RGBColor.from_string(cfg["color_primary"])
        if rep.get("title_en"):
            en = self.doc.add_paragraph()
            er = en.add_run(_nbsp(rep["title_en"]))
            _set_run_font(er, cfg["font_heading"], 10.5)
            er.italic = True
            er.font.color.rgb = RGBColor(0x44, 0x4F, 0x62)
        meta = []
        if rep.get("edition_label"):
            meta.append(rep["edition_label"])
        if rep.get("publish_date"):
            meta.append(_fmt_date(rep["publish_date"]))
        meta.append("임시 간결형 요약본(IITP 공식 양식 아님)")
        p = self.doc.add_paragraph()
        mr = p.add_run(" · ".join(meta))
        _set_run_font(mr, cfg["font_body"], 9)
        mr.font.color.rgb = RGBColor(0x55, 0x60, 0x70)
        ppr = p._p.get_or_add_pPr()
        ppr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="2" w:color="{cfg["color_primary"]}"/></w:pBdr>'))

    def _page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- 본문 ----
    def build(self) -> BuildResult:
        self._setup_styles()
        toc_anchor = None
        if self.compact:
            self._build_compact_title()
        else:
            self._build_cover()
            toc_title = self.doc.add_paragraph()
            tr = toc_title.add_run("목  차")
            _set_run_font(tr, self.config["font_heading"], 17)
            tr.bold = True
            tr.font.color.rgb = RGBColor.from_string(self.config["color_primary"])
            toc_anchor = self.doc.add_paragraph()  # 본문을 다 만든 뒤 이 앞에 목차 항목을 삽입한다
            self._page_break()

        for part_no, part in enumerate(self.snapshot["parts"], start=1):
            self._build_part(part_no, part)

        if toc_anchor is not None:
            self._insert_toc(toc_anchor)
        self._attach_footnotes_part()
        buf = io.BytesIO()
        self.doc.save(buf)
        return BuildResult(
            docx_bytes=buf.getvalue(), footnote_count=len(self.footnotes), table_count=self.table_count,
            headings=list(self.headings), part_headings=list(self.part_headings),
        )

    def _bookmark(self, paragraph, idx: int) -> None:
        """목차 항목이 제목 문자열이 아니라 '몇 번째 제목인가'로 연결되도록 고유 북마크(_Toc000001…)를 건다."""
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(idx + 1))
        start.set(qn("w:name"), f"_Toc{idx + 1:06d}")
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(idx + 1))
        p = paragraph._p
        ppr = p.find(qn("w:pPr"))
        if ppr is not None:
            ppr.addnext(start)
        else:
            p.insert(0, start)
        p.append(end)

    def _add_heading(self, text: str, level: int) -> None:
        p = self.doc.add_heading(level=level)
        run = p.add_run(_nbsp(text))
        _set_run_font(run, self.config["font_heading"])
        if level <= 3:
            self._bookmark(p, len(self.headings))
        self.headings.append((level, text))

    def _build_part(self, part_no: int, part: dict) -> None:
        self._current_part_no = part_no
        self._table_counter_in_part = 0
        has_content = bool(part.get("manuscripts")) or bool(part.get("summary_items"))
        if part_no > 1 and has_content and not self.compact:  # 내용이 있는 파트만 새 페이지에서 시작(빈 파트로 빈 페이지가 생기지 않게)
            self._page_break()
        self._add_heading(part["heading"], 1)
        self.part_headings.append(part["heading"])

        manuscripts = part.get("manuscripts", [])
        if not manuscripts and not part.get("summary_items"):
            p = self.doc.add_paragraph()
            r = p.add_run("(원고 미배치 — 미완성 초안입니다. 작성 후 다시 출력하세요.)")
            _set_run_font(r, self.config["font_body"], 10)
            r.italic = True
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
            return

        for ms in manuscripts:
            self._fn_lookup = ms.get("footnotes", {})
            if not ms.get("hide_title"):
                self._add_heading(ms["title"], 2)
            if ms.get("title_en"):
                p = self.doc.add_paragraph()
                r = p.add_run(ms["title_en"])
                _set_run_font(r, self.config["font_body"], 10)
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x60, 0x70)
            if ms.get("byline"):  # 주간기술동향지: 저자(소속) 표시
                p = self.doc.add_paragraph()
                r = p.add_run(ms["byline"])
                _set_run_font(r, self.config["font_body"], 10.5)
                r.bold = True
                r.font.color.rgb = RGBColor(0x44, 0x4F, 0x62)
            if ms.get("summary"):
                p = self.doc.add_paragraph()
                r = p.add_run("요약  " + ms["summary"])
                _set_run_font(r, self.config["font_body"], 10)
                r.italic = True
            if ms.get("status_note"):
                p = self.doc.add_paragraph()
                r = p.add_run(ms["status_note"])
                _set_run_font(r, self.config["font_body"], 8.5)
                r.font.color.rgb = RGBColor(0x9A, 0x4B, 0x00)
            self._render_manuscript_body(ms)

        for item in part.get("summary_items", []):
            self._fn_lookup = item.get("footnotes", {})
            self._render_summary_item(item)

    def _render_manuscript_body(self, ms: dict) -> None:
        body = ms.get("body_json") or {"type": "doc", "content": []}
        blocks = list(body.get("content", []))
        # 편집기 본문 첫 제목이 원고 제목과 같으면 중복이므로 건너뛴다
        if blocks and blocks[0].get("type") == "heading":
            first_text = "".join(n.get("text", "") for n in blocks[0].get("content", []))
            a, b = _title_key(first_text), _title_key(ms["title"])
            if a and b and (a == b or a in b or b in a):
                blocks = blocks[1:]
        levels = [b["attrs"].get("level", 1) for b in blocks if b.get("type") == "heading"]
        base = min(levels) if levels else 1
        self._heading_shift = 3 - base  # 본문에서 가장 높은 제목이 Word '제목 3'(목차 3단계)이 되도록 맞춘다
        self._render_blocks(self.doc, blocks)

    def _render_summary_item(self, item: dict) -> None:
        p = self.doc.add_paragraph(style="List Bullet")
        self._add_inline_text(p, item["text"])
        for idx, fn_id in enumerate(item.get("footnote_ids", [])):
            if idx > 0:  # 각주 번호가 연달아 붙어 '12'처럼 보이지 않도록 위첨자 쉼표로 구분
                sep = p.add_run(",")
                sep.font.superscript = True
                _set_run_font(sep, self.config["font_body"], self.config["font_size_body"])
            self._add_footnote_ref(p, fn_id)
        if item.get("needs_review"):
            r = p.add_run("  [원고 변경 — 요약문 재검토 필요]")
            _set_run_font(r, self.config["font_body"], 8.5)
            r.font.color.rgb = RGBColor(0x9A, 0x4B, 0x00)

    # ---- ProseMirror 블록 → DOCX ----
    def _render_blocks(self, container, blocks: list[dict], list_depth: int = 0, first_par=None) -> None:
        use_first = first_par
        for block in blocks:
            btype = block.get("type")
            if btype == "paragraph":
                par = use_first if use_first is not None else container.add_paragraph()
                use_first = None
                self._add_inline(par, block.get("content", []))
            elif btype == "heading":
                level = int(block.get("attrs", {}).get("level", 1))
                word_level = max(1, min(4, level + getattr(self, "_heading_shift", 2)))
                text = "".join(n.get("text", "") for n in block.get("content", []))
                p = self.doc.add_heading(level=word_level)
                self._add_inline(p, block.get("content", []), font_override=self.config["font_heading"])
                if word_level <= 3 and container is self.doc:
                    self._bookmark(p, len(self.headings))
                    self.headings.append((word_level, text))
            elif btype in ("bulletList", "orderedList"):
                self._render_list(container, block, list_depth)
            elif btype == "blockquote":
                for child in block.get("content", []):
                    par = container.add_paragraph()
                    ppr = par._p.get_or_add_pPr()
                    ppr.append(parse_xml(
                        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="6" w:color="9AA7BD"/></w:pBdr>'
                    ))
                    par.paragraph_format.left_indent = Mm(6)  # 스키마 순서상 pBdr 뒤에 ind가 와야 한다
                    self._add_inline(par, child.get("content", []))
            elif btype == "table":
                self._render_table(container, block)
            elif btype == "horizontalRule":
                par = container.add_paragraph()
                par._p.get_or_add_pPr().append(parse_xml(
                    f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="9AA7BD"/></w:pBdr>'
                ))
            elif btype == "codeBlock":
                par = container.add_paragraph()
                text = "".join(n.get("text", "") for n in block.get("content", []))
                r = par.add_run(text)
                _set_run_font(r, "Consolas", 9)
            # 알 수 없는 노드는 조용히 버리지 않고 평문으로 남긴다
            elif block.get("content"):
                par = container.add_paragraph()
                self._add_inline(par, block.get("content", []))

    def _render_list(self, container, block: dict, depth: int) -> None:
        ordered = block["type"] == "orderedList"
        for idx, li in enumerate(block.get("content", []), start=1):
            children = li.get("content", [])
            para_children = [c for c in children if c.get("type") == "paragraph"]
            nested = [c for c in children if c.get("type") in ("bulletList", "orderedList")]
            style = "List Bullet 2" if depth >= 1 else "List Bullet"
            par = container.add_paragraph(style=None if ordered else style) if container is not self.doc else (
                self.doc.add_paragraph(style=None if ordered else style))
            if ordered:
                # Word 번호 목록은 문서 전체에서 번호가 이어지는 문제가 있어 목록별로 1부터 다시 시작하도록 수동 번호를 쓴다
                par.paragraph_format.left_indent = Mm(6 + 5 * depth)
                par.paragraph_format.first_line_indent = Mm(-5)
                r = par.add_run(f"{idx}. ")
                _set_run_font(r, self.config["font_body"], self.config["font_size_body"])
            if para_children:
                self._add_inline(par, para_children[0].get("content", []))
                for extra in para_children[1:]:
                    p2 = container.add_paragraph()
                    p2.paragraph_format.left_indent = Mm(6 + 5 * depth)
                    self._add_inline(p2, extra.get("content", []))
            for sub in nested:
                self._render_list(container, sub, depth + 1)

    def _render_table(self, container, block: dict) -> None:
        rows = block.get("content", [])
        if not rows:
            return
        self.table_count += 1
        self._table_counter_in_part += 1
        caption_text = self.config["table_caption_format"].format(part=self._current_part_no, n=self._table_counter_in_part)
        user_caption = (block.get("attrs", {}) or {}).get("caption", "").strip()
        cap = container.add_paragraph(style="Caption")
        cap.paragraph_format.keep_with_next = True
        cap.add_run(caption_text + (f" {user_caption}" if user_caption else ""))  # [표 1-1] 국가별 주요 정책 비교

        n_cols = 0
        for row in rows:
            width = sum(int(c.get("attrs", {}).get("colspan", 1)) for c in row.get("content", []))
            n_cols = max(n_cols, width)
        table = container.add_table(rows=len(rows), cols=max(n_cols, 1))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tblpr = table._tbl.tblPr
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is not None:  # 스키마 순서를 지키기 위해 기존 요소의 값만 바꾼다
            tblw.set(qn("w:w"), "5000")
            tblw.set(qn("w:type"), "pct")

        occupied: set[tuple[int, int]] = set()
        for r_idx, row in enumerate(rows):
            col = 0
            for cell_node in row.get("content", []):
                while (r_idx, col) in occupied:
                    col += 1
                attrs = cell_node.get("attrs", {})
                colspan, rowspan = int(attrs.get("colspan", 1)), int(attrs.get("rowspan", 1))
                cell = table.cell(r_idx, col)
                target = cell
                if colspan > 1 or rowspan > 1:
                    target = cell.merge(table.cell(min(r_idx + rowspan - 1, len(rows) - 1), min(col + colspan - 1, n_cols - 1)))
                for rr in range(r_idx, r_idx + rowspan):
                    for cc in range(col, col + colspan):
                        occupied.add((rr, cc))
                is_header = cell_node.get("type") == "tableHeader"
                content = cell_node.get("content", [])
                first = target.paragraphs[0]
                self._render_blocks(target, content, first_par=first)
                if is_header:
                    self._shade_cell(target, self.config["color_table_header"])
                    for p in target.paragraphs:
                        for run in p.runs:
                            run.bold = True
                for p in target.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    p.paragraph_format.line_spacing = 1.15
                    for run in p.runs:
                        if run.font.size is None:
                            run.font.size = Pt(9.5)
                col += colspan
            # 제목행(첫 행이 tableHeader로만 구성된 경우) 반복 + 행이 쪼개지지 않게 한다
            if r_idx == 0:  # 표 제목행과 첫 데이터 행이 쪽 경계에서 떨어지지 않게 한다
                for c in table.rows[0].cells:
                    for p in c.paragraphs:
                        p.paragraph_format.keep_with_next = True
            trpr = table.rows[r_idx]._tr.get_or_add_trPr()
            trpr.append(OxmlElement("w:cantSplit"))
            if r_idx == 0 and all(c.get("type") == "tableHeader" for c in row.get("content", [])):
                trpr.append(OxmlElement("w:tblHeader"))
        spacer = container.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)

    @staticmethod
    def _shade_cell(cell, hex_color: str) -> None:
        tcpr = cell._tc.get_or_add_tcPr()
        tcpr.append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{hex_color}"/>'))

    # ---- 인라인(텍스트·마크·각주·링크) ----
    def _add_inline_text(self, par, text: str) -> None:
        r = par.add_run(text)
        _set_run_font(r, self.config["font_body"], self.config["font_size_body"])

    def _add_inline(self, par, nodes: list[dict], font_override: str | None = None) -> None:
        pending: str | None = None
        for node in nodes:
            ntype = node.get("type")
            fn_id = None
            for mark in node.get("marks", []) or []:
                if mark.get("type") == "footnoteMark":
                    fn_id = mark.get("attrs", {}).get("footnoteId")
            if pending and fn_id != pending:
                self._add_footnote_ref(par, pending)
                pending = None
            if ntype == "hardBreak":
                par.add_run().add_break()
            elif ntype == "text":
                self._add_text_node(par, node, font_override)
            if fn_id:
                pending = fn_id
        if pending:
            self._add_footnote_ref(par, pending)

    def _add_text_node(self, par, node: dict, font_override: str | None) -> None:
        text = node.get("text", "")
        if not text:
            return
        marks = {m.get("type"): m.get("attrs", {}) for m in node.get("marks", []) or []}
        font = font_override or self.config["font_body"]
        if "link" in marks and marks["link"].get("href"):
            self._add_hyperlink(par, marks["link"]["href"], text, font)
            return
        r = par.add_run(text)
        _set_run_font(r, font, None if font_override else self.config["font_size_body"])
        if "bold" in marks:
            r.bold = True
        if "italic" in marks:
            r.italic = True
        if "underline" in marks:
            r.underline = True
        if "strike" in marks:
            r.font.strike = True
        if "code" in marks:
            _set_run_font(r, "Consolas", 9.5)

    def _add_hyperlink(self, par, url: str, text: str, font: str) -> None:
        r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
        hl = OxmlElement("w:hyperlink")
        hl.set(qn("r:id"), r_id)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        rstyle = OxmlElement("w:rStyle")
        rstyle.set(qn("w:val"), "Hyperlink")
        rpr.append(rstyle)
        rfonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), font)
        rpr.append(rfonts)
        run.append(rpr)
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        run.append(t)
        hl.append(run)
        par._p.append(hl)

    def _add_footnote_ref(self, par, mark_id: str) -> None:
        fn = self._fn_lookup.get(mark_id)
        if fn is None:
            return
        fn_id = len(self.footnotes) + 1
        source_key = fn.get("source_url") or fn.get("doc_title") or mark_id
        first = self._seen_sources.get(source_key)
        if first is None:
            self._seen_sources[source_key] = fn_id
            self.footnotes.append((fn_id, format_footnote_text(fn, self.config), fn.get("source_url")))
        else:  # 같은 출처를 다시 인용하면 전체 서지정보를 반복하지 않고 첫 각주를 가리킨다(근거 검증 상태 표시는 각주마다 유지)
            self.footnotes.append((fn_id, format_footnote_text(fn, self.config, short_ref=first), None))
        run = par.add_run()
        rpr = run._r.get_or_add_rPr()
        rstyle = OxmlElement("w:rStyle")
        rstyle.set(qn("w:val"), "FootnoteReference")
        rpr.insert(0, rstyle)
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(fn_id))
        run._r.append(ref)

    # ---- 목차 ----
    def _insert_toc(self, anchor) -> None:
        entries = [h for h in self.headings if h[0] <= 3]
        cfg = self.config
        known_pages = self.toc_pages is not None
        paragraphs = []
        for idx, (level, text) in enumerate(entries):
            p = OxmlElement("w:p")
            ppr = OxmlElement("w:pPr")
            tabs = parse_xml(
                f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="9000"/></w:tabs>'
            )
            ppr.append(tabs)
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str((level - 1) * 300))
            ppr.append(ind)
            p.append(ppr)
            if idx == 0:
                for kind, instr in (("begin", None), (None, 'TOC \\o "1-3" \\h \\z \\u'), ("separate", None)):
                    r = OxmlElement("w:r")
                    if kind:
                        fc = OxmlElement("w:fldChar")
                        fc.set(qn("w:fldCharType"), kind)
                        if kind == "begin" and not known_pages:
                            fc.set(qn("w:dirty"), "true")  # 페이지 번호를 모를 때는 Word가 열 때 갱신하도록 표시
                        r.append(fc)
                    else:
                        it = OxmlElement("w:instrText")
                        it.set(qn("xml:space"), "preserve")
                        it.text = f" {instr} "
                        r.append(it)
                    p.append(r)
            page = self.toc_pages[idx] if known_pages and idx < len(self.toc_pages) else None
            line = text + ("\t" + str(page) if page else "")
            hl = OxmlElement("w:hyperlink")  # 목차 항목을 누르면 해당 제목으로 이동(Word·PDF)
            hl.set(qn("w:anchor"), f"_Toc{idx + 1:06d}")
            hl.set(qn("w:history"), "1")
            r = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                rfonts.set(qn(attr), cfg["font_body"])
            rpr.append(rfonts)
            if level == 1:
                rpr.append(OxmlElement("w:b"))
            r.append(rpr)
            if "\t" in line:
                t1 = OxmlElement("w:t")
                t1.set(qn("xml:space"), "preserve")
                t1.text = line.split("\t")[0]
                r.append(t1)
                r.append(OxmlElement("w:tab"))
                t2 = OxmlElement("w:t")
                t2.text = line.split("\t")[1]
                r.append(t2)
            else:
                t = OxmlElement("w:t")
                t.set(qn("xml:space"), "preserve")
                t.text = line
                r.append(t)
            hl.append(r)
            p.append(hl)
            if idx == len(entries) - 1:
                r_end = OxmlElement("w:r")
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), "end")
                r_end.append(fc)
                p.append(r_end)
            paragraphs.append(p)
        for p in paragraphs:
            anchor._p.addprevious(p)
        if not known_pages:
            note = self.doc.add_paragraph()
            note._p.getparent().remove(note._p)
            n = OxmlElement("w:p")
            run = OxmlElement("w:r")
            rpr = OxmlElement("w:rPr")
            rpr.append(parse_xml(f'<w:color {nsdecls("w")} w:val="888888"/>'))
            rpr.append(parse_xml(f'<w:sz {nsdecls("w")} w:val="16"/>'))
            run.append(rpr)
            t = OxmlElement("w:t")
            t.text = "※ 목차의 페이지 번호는 Word에서 목차를 우클릭 → '필드 업데이트'로 갱신하세요(이 파일 생성 시점에는 렌더링으로 확인하지 않아 번호를 넣지 않았습니다)."
            run.append(t)
            n.append(run)
            anchor._p.addprevious(n)

    # ---- 각주 파트 ----
    def _attach_footnotes_part(self) -> None:
        w_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        package = self.doc.part.package
        part = Part(
            PackURI("/word/footnotes.xml"),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
            b"", package,
        )
        self.doc.part.relate_to(part, RT.FOOTNOTES)
        font = self.config["font_body"]
        fonts = f'<w:rFonts w:ascii="{font}" w:hAnsi="{font}" w:eastAsia="{font}" w:cs="{font}"/>'

        def esc(t: str) -> str:
            return t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        def run(t: str) -> str:
            return f'<w:r><w:rPr>{fonts}</w:rPr><w:t xml:space="preserve">{esc(t)}</w:t></w:r>' if t else ""

        parts = [
            f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{w_ns}" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">',
            '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
            "<w:r><w:separator/></w:r></w:p></w:footnote>",
            '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
            "<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>",
        ]
        for fn_id, text, url in self.footnotes:
            if url and url in text:  # 각주 속 URL을 Word에서 클릭 가능한 실제 하이퍼링크로 만든다
                head, tail = text.split(url, 1)
                rid = part.relate_to(url, RT.HYPERLINK, is_external=True)
                body = (run(" " + head)
                        + f'<w:hyperlink r:id="{rid}"><w:r><w:rPr><w:rStyle w:val="Hyperlink"/>{fonts}</w:rPr>'
                          f'<w:t xml:space="preserve">{esc(url)}</w:t></w:r></w:hyperlink>' + run(tail))
            else:
                body = run(" " + text)
            parts.append(
                f'<w:footnote w:id="{fn_id}"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
                f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>{body}</w:p></w:footnote>'
            )
        parts.append("</w:footnotes>")
        part._blob = "".join(parts).encode("utf-8")


def build_docx(snapshot: dict, toc_pages: list[int | None] | None = None) -> BuildResult:
    return DocxBuilder(snapshot, toc_pages).build()
