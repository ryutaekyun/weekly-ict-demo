"""Phase 3-2 — 보고서 구성·발간 점검·DOCX/PDF 출력·확정·권한 테스트.

생성된 DOCX/PDF는 실제로 열어서(zip·python-docx·PyMuPDF) 내용을 확인한다. PDF는 LibreOffice가 있는 환경에서만
실행하고(없으면 skip), CI(Ubuntu)에서는 반드시 설치해 실제 변환을 검증한다.
"""

import io
import json
import os
import uuid
import zipfile

import pytest
from docx import Document

from app.core.config import settings
from app.models.enums import UserRole
from app.models.report import ReportExport
from app.services import report_docx, report_pdf
from tests.conftest import auth_headers, make_user

_has_soffice = report_pdf.find_soffice() is not None
requires_soffice = pytest.mark.skipif(not _has_soffice, reason="LibreOffice 없음 — CI에서 검증")


@pytest.fixture(autouse=True)
def _export_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "export_dir", str(tmp_path / "exports"))


@pytest.fixture(autouse=True)
def _skip_rendering_unless_marked(request, monkeypatch):
    """LibreOffice 변환은 건당 수십 초가 걸려, 실제 렌더링 검증이 목적인 테스트(@pytest.mark.render)에서만 켠다.
    나머지 테스트는 DOCX 생성·권한·확정 로직만 빠르게 검증한다."""
    if request.node.get_closest_marker("render") is None:
        monkeypatch.setattr(report_pdf, "find_soffice", lambda: None)


def test_libreoffice_is_available_when_required_by_ci():
    """CI에서는 PDF 렌더링 검증이 생략(skip)되어 통과로 보이는 일이 없도록 LibreOffice 설치를 필수로 확인한다."""
    if os.environ.get("REQUIRE_SOFFICE") == "1":
        assert report_pdf.find_soffice() is not None, "CI에 LibreOffice가 설치되지 않았습니다"


def _t(text, fn=None, marks=None):
    node = {"type": "text", "text": text}
    all_marks = list(marks or [])
    if fn:
        all_marks.append({"type": "footnoteMark", "attrs": {"footnoteId": fn}})
    if all_marks:
        node["marks"] = all_marks
    return node


def _cell(kind, text):
    return {"type": kind, "attrs": {"colspan": 1, "rowspan": 1}, "content": [{"type": "paragraph", "content": [_t(text)]}]}


def _body_doc(title, mark_id):
    return {"type": "doc", "content": [
        {"type": "heading", "attrs": {"level": 1}, "content": [_t(title)]},
        {"type": "heading", "attrs": {"level": 2}, "content": [_t("1. 추진 경과")]},
        {"type": "paragraph", "content": [_t("미국 NHTSA는 제정안을 공고함", mark_id), _t(". 시행일은 2025.2.18.임")]},
        {"type": "paragraph", "content": [_t("굵은 글씨", marks=[{"type": "bold"}]), _t(" 와 "),
                                          _t("링크", marks=[{"type": "link", "attrs": {"href": "https://example.org/x"}}])]},
        {"type": "bulletList", "content": [{"type": "listItem", "content": [{"type": "paragraph", "content": [_t("글머리 항목")]}]}]},
        {"type": "table", "content": [
            {"type": "tableRow", "content": [_cell("tableHeader", "구분"), _cell("tableHeader", "일정")]},
            {"type": "tableRow", "content": [_cell("tableCell", "시행일"), _cell("tableCell", "2025.2.18.")]},
        ]},
    ]}


def _plain(doc):
    lines = []

    def walk(n):
        if n.get("type") == "text":
            lines.append(n["text"])
        for c in n.get("content", []):
            walk(c)
        if n.get("type") in ("paragraph", "heading"):
            lines.append("\n")

    walk(doc)
    return "".join(lines).strip()


def _make_manuscript(client, headers, *, title="시범 원고", with_footnote=True, verify=False):
    created = client.post("/manuscripts", json={"manuscript_type": "key_issue", "title": title}, headers=headers)
    assert created.status_code == 201, created.text
    mid = created.json()["id"]
    mark_id = str(uuid.uuid4())
    doc = _body_doc(title, mark_id if with_footnote else "none")
    saved = client.patch(f"/manuscripts/{mid}", headers=headers, json={
        "body_json": json.dumps(doc, ensure_ascii=False), "body": _plain(doc), "lock_version": 0,
        "is_milestone": True, "change_note": "테스트",
    })
    assert saved.status_code == 200, saved.text
    if with_footnote:
        fn = client.post(f"/manuscripts/{mid}/footnotes", headers=headers, json={
            "mark_id": mark_id, "anchor_text": "미국 NHTSA는 제정안을 공고함",
            "source_url": "https://www.federalregister.gov/documents/test", "source_org": "NHTSA",
        })
        assert fn.status_code == 201, fn.text
        if verify:
            v = client.post(f"/manuscripts/{mid}/footnotes/{fn.json()['id']}/verify", headers=headers,
                            params={"matched": True})
            assert v.status_code == 200, v.text
    return mid


def _general_report(client, headers, manuscript_ids):
    created = client.post("/reports", json={"report_type": "general", "title": "테스트 일반호", "title_en": "Test"},
                          headers=headers)
    assert created.status_code == 201, created.text
    rid = created.json()["id"]
    for part, mid in manuscript_ids.items():
        r = client.post(f"/reports/{rid}/items", json={"part_key": part, "manuscript_id": mid}, headers=headers)
        assert r.status_code == 201, r.text
    return rid


# ---------------------------------------------------------------- DOCX 엔진 ----
def test_docx_engine_has_real_word_footnotes_table_header_and_hyperlink():
    fn = {"source_org": "NHTSA", "doc_title": "FMVSS 305a", "source_url": "https://www.federalregister.gov/x",
          "published_at": "2024-04-15T00:00:00+00:00", "accessed_at": "2026-10-23T00:00:00+00:00",
          "verification_status": "not_verified", "location_label": None}
    snap = {
        "report": {"title": "엔진 테스트", "title_en": "Engine", "kind_label": "일반호", "iitp_confirmed": False, "is_pilot": True},
        "template": {"name": "임시 표준 템플릿", "version": 1, "is_provisional": True, "config": {}},
        "parts": [{"key": "key_issue", "heading": "Ⅰ. 주요 이슈", "manuscripts": [
            {"title": "시범 원고", "title_en": None, "body_json": _body_doc("시범 원고", "m1"), "footnotes": {"m1": fn}}]},
            {"key": "country_trend", "heading": "Ⅱ. 주요국 동향", "manuscripts": []}],
    }
    built = report_docx.build_docx(snap)
    assert built.footnote_count == 1 and built.table_count == 1

    zf = zipfile.ZipFile(io.BytesIO(built.docx_bytes))
    assert zf.testzip() is None
    footnotes_xml = zf.read("word/footnotes.xml").decode("utf-8")
    assert "FMVSS 305a" in footnotes_xml and "[근거 미검증]" in footnotes_xml and "위치 미확인" in footnotes_xml
    document_xml = zf.read("word/document.xml").decode("utf-8")
    assert document_xml.count("<w:footnoteReference") == 1  # 실제 Word 각주 참조
    assert "w:tblHeader" in document_xml  # 표 제목행 반복
    assert "[표 1-1]" in document_xml
    assert "w:hyperlink" in document_xml
    assert "(원고 미배치" in document_xml  # 빈 파트는 임의로 채우지 않는다
    assert 'TOC \\o "1-3"' in document_xml
    # 렌더링으로 페이지를 확인하지 않았으므로 목차에 임의의 번호가 없고 안내 문구가 있어야 한다
    assert "페이지 번호는 Word에서" in document_xml

    reopened = Document(io.BytesIO(built.docx_bytes))  # python-docx로 다시 열 수 있어야(=수정 가능한 구조)
    assert any("추진 경과" in p.text for p in reopened.paragraphs)


def test_footnote_format_is_configurable_and_never_marks_unverified_as_verified():
    base = {"source_org": "기관", "doc_title": "자료", "source_url": "https://e.org", "published_at": None,
            "accessed_at": None, "verification_status": "not_verified", "location_label": "3문단"}
    cfg = report_docx.merge_template_config({"footnote_format": "{org} - {title} <{url}>"})
    text = report_docx.format_footnote_text(base, cfg)
    assert text.startswith("기관 - 자료 <https://e.org>") and "[근거 미검증]" in text and "위치 미확인" not in text
    matched = report_docx.format_footnote_text({**base, "verification_status": "matched"}, cfg)
    assert "[근거 미검증]" not in matched


# ---------------------------------------------------------------- API ----
def test_report_endpoints_require_auth(client):
    assert client.get("/reports").status_code == 401
    assert client.post("/reports", json={"report_type": "general", "title": "x"}).status_code == 401


def test_general_report_preflight_flags_missing_parts_and_unverified(client, db_session):
    make_user(db_session, "r1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r1@example.com")
    mid = _make_manuscript(client, h)
    rid = _general_report(client, h, {"key_issue": mid})
    detail = client.get(f"/reports/{rid}", headers=h).json()
    codes = {i["code"] for i in detail["preflight"]}
    assert "required_part_missing" in codes  # 주요국 동향 없음
    assert "unverified_footnotes" in codes
    assert "iitp_unconfirmed" in codes
    assert len(detail["items"]) == 1


def test_docx_export_creates_real_file_with_hash_and_snapshot(client, db_session, tmp_path):
    make_user(db_session, "r2@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r2@example.com")
    mid = _make_manuscript(client, h)
    rid = _general_report(client, h, {"key_issue": mid})
    resp = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h)
    assert resp.status_code == 201, resp.text
    exp = resp.json()
    assert exp["status"] == "success", exp
    assert exp["footnote_count"] == 1 and exp["table_count"] == 1 and exp["file_size"] > 5000
    versions = json.loads(exp["manuscript_versions_json"])
    assert versions[mid]["lock_version"] == 1  # 출력에 사용한 원고 버전이 고정 기록된다

    dl = client.get(f"/reports/{rid}/exports/{exp['id']}/download", headers=h)
    assert dl.status_code == 200
    import hashlib

    assert hashlib.sha256(dl.content).hexdigest() == exp["file_sha256"]
    doc = Document(io.BytesIO(dl.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "시범 원고" in text and "추진 경과" in text and "Ⅰ. 주요 이슈" in text
    assert b"[\xea\xb7\xbc\xea\xb1\xb0 \xeb\xaf\xb8\xea\xb2\x80\xec\xa6\x9d]" in zipfile.ZipFile(io.BytesIO(dl.content)).read("word/footnotes.xml")

    listed = client.get(f"/reports/{rid}/exports", headers=h).json()
    assert len(listed) == 1


@requires_soffice
@pytest.mark.render
def test_pdf_export_renders_real_pages_with_korean_footnotes_and_toc_numbers(client, db_session):
    make_user(db_session, "r3@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r3@example.com")
    mid = _make_manuscript(client, h)
    rid = _general_report(client, h, {"key_issue": mid})
    resp = client.post(f"/reports/{rid}/exports", json={"file_format": "pdf"}, headers=h)
    assert resp.status_code == 201, resp.text
    exp = resp.json()
    assert exp["status"] == "success", exp
    assert exp["page_count"] >= 3
    dl = client.get(f"/reports/{rid}/exports/{exp['id']}/download", headers=h)
    assert dl.content[:4] == b"%PDF"
    analysis = report_pdf.analyze_pdf(dl.content, [(1, "Ⅰ. 주요 이슈")], ["Ⅰ. 주요 이슈"])
    text = analysis["text"]
    assert "주요 이슈" in text and "추진 경과" in text  # 한글이 깨지지 않고 추출된다
    assert "NHTSA" in text and "근거 미검증" in text  # 각주가 페이지 하단에 실제 출력됨
    assert "목" in text and "차" in text
    validation = json.loads(exp["validation_json"])
    assert validation["rendered"] is True and validation["part_pages"][0]["pages"] >= 1

    sample_dir = os.environ.get("REPORT_SAMPLE_DIR")  # CI가 합성 데이터 출력물(PDF·페이지 이미지)을 아티팩트로 보관할 때 사용
    if sample_dir:
        from pathlib import Path

        out = Path(sample_dir)
        out.mkdir(parents=True, exist_ok=True)
        (out / "ci_sample.pdf").write_bytes(dl.content)
        report_pdf.render_pages_png(dl.content, out / "pages")


@requires_soffice
@pytest.mark.render
def test_docx_toc_contains_real_rendered_page_numbers_when_libreoffice_available(client, db_session):
    make_user(db_session, "r7@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r7@example.com")
    mid = _make_manuscript(client, h)
    rid = _general_report(client, h, {"key_issue": mid})
    exp = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h).json()
    assert exp["status"] == "success" and exp["page_count"] >= 3
    dl = client.get(f"/reports/{rid}/exports/{exp['id']}/download", headers=h)
    xml = zipfile.ZipFile(io.BytesIO(dl.content)).read("word/document.xml").decode("utf-8")
    assert "페이지 번호는 Word에서" not in xml  # 렌더링으로 확인한 번호가 들어갔으므로 안내 문구가 없다
    assert 'w:dirty="true"' not in xml


def test_summary_report_edit_persist_and_stale_marking(client, db_session):
    make_user(db_session, "r4@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r4@example.com")
    mid = _make_manuscript(client, h)
    sentences = client.get(f"/reports/source-sentences/{mid}", headers=h).json()["sentences"]
    source = next(s for s in sentences if "제정안" in s)
    created = client.post("/reports", json={"report_type": "summary", "title": "요약본"}, headers=h).json()
    rid = created["id"]
    added = client.post(f"/reports/{rid}/summary-items", headers=h,
                        json={"part_key": "key_issue", "manuscript_id": mid, "source_sentence": source})
    assert added.status_code == 201, added.text
    item = added.json()["items"][0]
    assert item["summary_text"] == source and item["needs_review"] is False

    edited = client.patch(f"/reports/{rid}/items/{item['id']}", headers=h, json={"summary_text": "연구원이 다듬은 요약문"})
    assert edited.json()["items"][0]["summary_text"] == "연구원이 다듬은 요약문"

    # 원고 본문에서 해당 문장이 바뀌면 '재검토 필요'가 표시되지만 연구원이 편집한 요약문은 덮어쓰이지 않는다
    ms = client.get(f"/manuscripts/{mid}", headers=h).json()
    changed = {"type": "doc", "content": [{"type": "paragraph", "content": [_t("완전히 다른 문장으로 바뀐 본문입니다")]}]}
    client.patch(f"/manuscripts/{mid}", headers=h, json={
        "body_json": json.dumps(changed, ensure_ascii=False), "body": _plain(changed), "lock_version": ms["lock_version"]})
    after = client.get(f"/reports/{rid}", headers=h).json()
    assert after["items"][0]["needs_review"] is True
    assert after["items"][0]["summary_text"] == "연구원이 다듬은 요약문"
    assert any(i["code"] == "summary_stale" for i in after["preflight"])

    bad = client.post(f"/reports/{rid}/summary-items", headers=h,
                      json={"part_key": "key_issue", "manuscript_id": mid, "source_sentence": "원고에 없는 문장입니다"})
    assert bad.status_code == 422


def test_finalize_blocked_until_checks_pass_or_approver_gives_reason(client, db_session):
    make_user(db_session, "author@example.com", UserRole.RESEARCHER)
    make_user(db_session, "boss@example.com", UserRole.APPROVER)
    ha = auth_headers(client, "author@example.com")
    hb = auth_headers(client, "boss@example.com")
    mid = _make_manuscript(client, ha)
    rid = _general_report(client, ha, {"key_issue": mid})

    assert client.post(f"/reports/{rid}/finalize", json={}, headers=ha).status_code == 403  # 연구원은 확정 불가
    blocked = client.post(f"/reports/{rid}/finalize", json={}, headers=hb)
    assert blocked.status_code == 422
    assert any(b["code"] == "required_part_missing" for b in blocked.json()["detail"]["blocking"])

    ok = client.post(f"/reports/{rid}/finalize", json={"exception_reason": "내부 점검용 임시 확정"}, headers=hb)
    assert ok.status_code == 200 and ok.json()["status"] == "final"
    assert ok.json()["finalize_exception_reason"] == "내부 점검용 임시 확정"
    assert ok.json()["iitp_confirmed"] is False  # 출력용 확정 ≠ IITP 공식 확인

    # 확정 후에는 작성자가 구성을 바꿀 수 없다
    mid2 = _make_manuscript(client, ha, title="다른 원고")
    edit = client.post(f"/reports/{rid}/items", json={"part_key": "country_trend", "manuscript_id": mid2}, headers=ha)
    assert edit.status_code == 409
    assert client.post(f"/reports/{rid}/reopen", headers=ha).status_code == 403
    assert client.post(f"/reports/{rid}/reopen", headers=hb).status_code == 200


def test_final_flow_passes_when_everything_is_verified_and_approved(client, db_session):
    make_user(db_session, "a2@example.com", UserRole.RESEARCHER)
    make_user(db_session, "boss2@example.com", UserRole.APPROVER)
    make_user(db_session, "rev2@example.com", UserRole.REVIEWER)
    ha, hb, hr = (auth_headers(client, e) for e in ("a2@example.com", "boss2@example.com", "rev2@example.com"))
    m1 = _make_manuscript(client, ha, title="주요 이슈 원고", verify=True)
    m2 = _make_manuscript(client, ha, title="주요국 동향 원고", verify=True)
    reviewer_id = client.get("/users", headers=ha).json()
    reviewer = next(u for u in reviewer_id if u["display_name"] == "테스트사용자" and u["role"] == "reviewer")
    for mid in (m1, m2):
        assert client.post(f"/manuscripts/{mid}/review-request", headers=ha, json={"reviewer_user_id": reviewer["id"]}).status_code == 200
        assert client.post(f"/manuscripts/{mid}/approve", headers=hb, json={}).status_code == 200
    rid = _general_report(client, ha, {"key_issue": m1, "country_trend": m2})
    final = client.post(f"/reports/{rid}/finalize", json={}, headers=hb)
    assert final.status_code == 200, final.text
    assert final.json()["blocking"] == []
    exp = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=ha).json()
    assert exp["is_final"] is True and exp["status"] == "success"
    assert "_확정" in exp["file_name"]


def test_other_users_cannot_download_or_list_files_and_duplicate_export_is_rejected(client, db_session):
    make_user(db_session, "owner@example.com", UserRole.RESEARCHER)
    make_user(db_session, "other@example.com", UserRole.RESEARCHER)
    ho = auth_headers(client, "owner@example.com")
    hx = auth_headers(client, "other@example.com")
    mid = _make_manuscript(client, ho)
    rid = _general_report(client, ho, {"key_issue": mid})
    exp = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=ho).json()

    assert client.get(f"/reports/{rid}/exports/{exp['id']}/download", headers=hx).status_code == 403
    assert client.get(f"/reports/{rid}/exports", headers=hx).status_code == 403
    assert client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=hx).status_code == 403
    assert client.post(f"/reports/{rid}/items", headers=hx, json={"part_key": "key_issue", "manuscript_id": mid}).status_code == 403
    client.cookies.clear()  # 테스트 클라이언트가 보관한 로그인 쿠키 제거 → 비로그인 상태
    assert client.get(f"/reports/{rid}/exports/{exp['id']}/download").status_code == 401

    # 진행 중인 출력이 있으면 같은 형식의 중복 요청은 거절한다
    db_session.add(ReportExport(report_id=uuid.UUID(rid), file_format="pdf", status="running",
                                created_by_user_id=db_session.query(__import__("app.models.user", fromlist=["User"]).User)
                                .filter_by(email="owner@example.com").one().id))
    db_session.commit()
    dup = client.post(f"/reports/{rid}/exports", json={"file_format": "pdf"}, headers=ho)
    assert dup.status_code == 409


def test_export_failure_is_recorded_and_retry_works(client, db_session, monkeypatch):
    make_user(db_session, "r5@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r5@example.com")
    mid = _make_manuscript(client, h)
    rid = _general_report(client, h, {"key_issue": mid})
    monkeypatch.setattr(report_pdf, "find_soffice", lambda: None)
    failed = client.post(f"/reports/{rid}/exports", json={"file_format": "pdf"}, headers=h)
    assert failed.status_code == 201
    assert failed.json()["status"] == "failed" and "LibreOffice" in failed.json()["error_message"]
    assert client.get(f"/reports/{rid}/exports/{failed.json()['id']}/download", headers=h).status_code == 404
    retry = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h)
    assert retry.json()["status"] == "success"  # LibreOffice 없이도 DOCX는 만들어지고, 렌더링 미확인이 명시된다
    assert "not_rendered" in retry.json()["validation_json"]


def test_template_new_version_keeps_old_version_and_admin_only(client, db_session):
    make_user(db_session, "adm@example.com", UserRole.ADMIN)
    make_user(db_session, "res@example.com", UserRole.RESEARCHER)
    ha = auth_headers(client, "adm@example.com")
    hr = auth_headers(client, "res@example.com")
    templates = client.get("/report-templates", headers=ha).json()
    assert templates[0]["version"] == 1 and templates[0]["is_provisional"] is True
    assert client.post(f"/report-templates/{templates[0]['id']}/new-version", headers=hr,
                       json={"config": {"footnote_format": "{org}"}}).status_code == 403
    new = client.post(f"/report-templates/{templates[0]['id']}/new-version", headers=ha,
                      json={"config": {"footnote_format": "{org} / {title}"}})
    assert new.status_code == 201 and new.json()["version"] == 2
    all_versions = client.get("/report-templates", headers=ha).json()
    assert {t["version"] for t in all_versions} == {1, 2}
    assert sum(1 for t in all_versions if t["is_active"]) == 1


def test_special_report_orders_manuscripts_and_allows_custom_headings(client, db_session):
    make_user(db_session, "r6@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r6@example.com")
    m1 = _make_manuscript(client, h, title="특집 A")
    m2 = _make_manuscript(client, h, title="특집 B")
    rid = client.post("/reports", json={"report_type": "special", "title": "특집호"}, headers=h).json()["id"]
    client.post(f"/reports/{rid}/items", headers=h, json={"manuscript_id": m1, "heading": "첫 번째 소주제"})
    detail = client.post(f"/reports/{rid}/items", headers=h, json={"manuscript_id": m2}).json()
    ids = [i["id"] for i in detail["items"]]
    reordered = client.post(f"/reports/{rid}/items/reorder", headers=h, json={"item_ids": list(reversed(ids))}).json()
    assert [i["manuscript_title"] for i in reordered["items"]] == ["특집 B", "특집 A"]
    exp = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h).json()
    assert exp["status"] == "success"
    dl = client.get(f"/reports/{rid}/exports/{exp['id']}/download", headers=h)
    paragraphs = [p.text for p in Document(io.BytesIO(dl.content)).paragraphs]
    assert paragraphs.index("1. 특집 B") < paragraphs.index("2. 첫 번째 소주제")
