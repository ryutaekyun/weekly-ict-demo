"""Phase 3-5 — 시연 기능 연결: 원문 링크 상태·검색 필터·후보주제 담당/이력/근거 원문·요약본 핵심 사실 변경 재검토·심의 예외."""

import io
import json
from datetime import datetime, timedelta, timezone

import httpx
import pytest
from docx import Document

from app.models.candidate import Candidate
from app.models.document_classification import DocumentClassification
from app.models.document_full_text import DocumentFullText
from app.models.enums import (
    CandidateStatus, ClassificationMethod, ClassificationRelevance, DocumentType, ExtractionMethod, ExtractionStatus, FieldCode,
    IssueMembershipType, SourceDocumentStatus, SourceType, UserRole,
)
from app.models.issue_cluster import IssueCluster
from app.models.issue_cluster_document import IssueClusterDocument
from app.models.source import Source
from app.models.source_document import SourceDocument
from app.services import journal_docx, link_quality
from tests.conftest import auth_headers, make_user


# ---------------------------------------------------------------- 원문 링크 분류·점검 ----
@pytest.mark.parametrize("url,state,blocked,pdf", [
    (None, "missing", True, False), ("", "missing", True, False), ("   ", "missing", True, False),
    ("ftp://example.org/a", "invalid", True, False), ("not a url", "invalid", True, False), ("https://", "invalid", True, False),
    ("https://www.example.org", "homepage", False, False), ("https://www.example.org/", "homepage", False, False),
    ("https://www.example.org/search?q=ai", "search", False, False), ("https://www.example.org/news?keyword=ai", "search", False, False),
    ("https://www.example.org/docs/2024/policy.pdf", "document", False, True), ("https://www.federalregister.gov/d/2025-02582", "document", False, False),
])
def test_classify_url(url, state, blocked, pdf):
    r = link_quality.classify_url(url)
    assert (r["state"], r["blocked"], r["is_pdf"]) == (state, blocked, pdf)
    assert (r["note"] is None) == (state == "document")  # 문제가 있으면 사유 문구가 있다


class _FakeClient:
    behavior = {"status": 200}

    def __init__(self, *a, **kw):
        pass

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def _resp(self, url):
        b = type(self).behavior
        if "raise" in b:
            raise b["raise"]
        return httpx.Response(b["status"], request=httpx.Request("GET", url))

    def head(self, url):
        return self._resp(url)

    def get(self, url):
        return self._resp(url)


@pytest.mark.parametrize("behavior,state", [
    ({"status": 200}, "reachable"), ({"status": 404}, "not_found"), ({"status": 410}, "not_found"), ({"status": 403}, "restricted"),
    ({"status": 429}, "restricted"), ({"status": 500}, "server_error"), ({"raise": httpx.ConnectTimeout("t")}, "unreachable"),
    ({"raise": httpx.ConnectError("c")}, "unreachable"),
])
def test_check_link_separates_deleted_url_from_restricted_access(monkeypatch, behavior, state):
    monkeypatch.setattr(link_quality.httpx, "Client", _FakeClient)
    _FakeClient.behavior = behavior
    r = link_quality.check_link("https://www.example.org/docs/1")
    assert r["state"] == state
    if state == "unreachable":
        assert "단정할 수 없습니다" in r["note"]  # 네트워크 제한을 주소 오류로 단정하지 않는다
    assert link_quality.check_link("")["state"] == "missing"  # 주소가 없으면 접속을 시도하지 않는다


# ---------------------------------------------------------------- 문서 목록: 분야·필터·정렬·건수 ----
def make_doc(db, *, title, field, url="https://example.org/doc", country="US", org="US Govt", extraction=ExtractionStatus.SUCCESS,
             days_ago=0, full_text="본문", relevance=ClassificationRelevance.MATCHED):
    src = db.query(Source).filter_by(org_name=org).first()
    if src is None:
        src = Source(name=f"{org} feed", org_name=org, field_code=FieldCode.SEMI, country_code=country, url="https://example.org/feed",
                     source_type=SourceType.API, is_primary_source=True)
        db.add(src)
        db.commit()
        db.refresh(src)
    published = datetime.now(timezone.utc) - timedelta(days=days_ago)
    doc = SourceDocument(source_id=src.id, title=title, origin_url=url, collected_at=datetime.now(timezone.utc), raw_text="",
                         content_hash=__import__("hashlib").sha256(f"{title}-{url}".encode()).hexdigest(), status=SourceDocumentStatus.COLLECTED, field_code=src.field_code,
                         collection_method=SourceType.API, document_type=DocumentType.POLICY, published_at=published, country_code=country)
    db.add(doc)
    db.commit()
    db.refresh(doc)
    db.add(DocumentClassification(document_id=doc.id, primary_field_code=field, relevance=relevance, method=ClassificationMethod.RULE_BASED,
                                  classified_at=datetime.now(timezone.utc)))
    if extraction is not None:
        db.add(DocumentFullText(document_id=doc.id, full_text=full_text, extraction_status=extraction, extraction_method=ExtractionMethod.HTML,
                                char_count=len(full_text), extracted_at=datetime.now(timezone.utc)))
    db.commit()
    return doc


def test_document_list_filters_sort_count_link_state_and_empty_field(client, db_session):
    make_user(db_session, "r@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r@example.com")
    make_doc(db_session, title="Alpha AI policy", field=FieldCode.AI, url="https://example.org/a", days_ago=3, full_text="신규 기술 설명 quantum")
    make_doc(db_session, title="Beta AI rule", field=FieldCode.AI, url="https://example.org/b.pdf", country="EU", org="EU Commission", days_ago=1,
             extraction=ExtractionStatus.FAILED, full_text="")
    make_doc(db_session, title="Gamma chip", field=FieldCode.SEMI, url="", days_ago=2)
    make_doc(db_session, title="Delta net", field=FieldCode.NET, url="https://www.example.org", days_ago=9, extraction=None)

    body = client.get("/documents", params={"field_code": "ai"}, headers=h).json()
    assert body["total"] == 2 and len(body["items"]) == 2  # 건수는 실제 조회된 데이터 수와 같다
    assert [d["title"] for d in body["items"]] == ["Beta AI rule", "Alpha AI policy"]  # 기본: 발표일 최신순
    assert client.get("/documents", params={"field_code": "ai", "sort": "published_asc"}, headers=h).json()["items"][0]["title"] == "Alpha AI policy"
    assert client.get("/documents", params={"sort": "title"}, headers=h).json()["items"][0]["title"] == "Alpha AI policy"
    assert client.get("/documents", params={"country_code": "EU"}, headers=h).json()["total"] == 1
    assert client.get("/documents", params={"extraction_status": "failed"}, headers=h).json()["total"] == 1
    assert client.get("/documents", params={"has_full_text": "true"}, headers=h).json()["total"] == 2  # 전문을 확보한 자료만
    assert client.get("/documents", params={"has_full_text": "false"}, headers=h).json()["total"] == 2
    assert client.get("/documents", params={"q": "quantum"}, headers=h).json()["total"] == 1  # 제목뿐 아니라 확보한 원문도 검색
    assert client.get("/documents", params={"q": "quantum", "field_code": "semi"}, headers=h).json()["total"] == 0
    items = {d["title"]: d for d in client.get("/documents", headers=h).json()["items"]}
    assert items["Gamma chip"]["link_blocked"] and items["Gamma chip"]["link_state"] == "missing" and items["Gamma chip"]["link_note"]
    assert items["Delta net"]["link_state"] == "homepage" and not items["Delta net"]["link_blocked"]
    assert items["Beta AI rule"]["is_pdf"] is True and items["Alpha AI policy"]["link_state"] == "document"
    # 자료가 없는 분야는 다른 분야 자료로 채우지 않고 빈 결과를 돌려준다
    empty = client.get("/documents", params={"field_code": "conv"}, headers=h).json()
    assert empty["total"] == 0 and empty["items"] == []
    facets = client.get("/documents/facets", headers=h).json()
    assert {c["code"] for c in facets["countries"]} == {"US", "EU"} and len(facets["sources"]) == 2 and facets["published_min"]


def test_check_link_endpoint_records_result_and_requires_login(client, db_session, monkeypatch):
    make_user(db_session, "r2@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r2@example.com")
    doc = make_doc(db_session, title="Link doc", field=FieldCode.AI, url="https://example.org/missing")
    monkeypatch.setattr(link_quality.httpx, "Client", _FakeClient)
    _FakeClient.behavior = {"status": 404}
    r = client.post(f"/documents/{doc.id}/check-link", headers=h).json()
    assert r["state"] == "not_found" and r["http_status"] == 404 and r["classification"]["state"] == "document"
    client.cookies.clear()
    assert client.post(f"/documents/{doc.id}/check-link").status_code == 401


# ---------------------------------------------------------------- 이슈·후보주제 연결, 담당자·이력 ----
def _cluster_with_docs(db, field=FieldCode.AI):
    d1 = make_doc(db, title="Issue doc one", field=field, url="https://example.org/i1", days_ago=5)
    d2 = make_doc(db, title="Issue doc two", field=field, url="", days_ago=2, org="EU Commission", country="EU")
    cluster = IssueCluster(title="Synthetic issue", field_code=field, first_seen_at=datetime.now(timezone.utc), last_activity_at=datetime.now(timezone.utc))
    db.add(cluster)
    db.commit()
    db.refresh(cluster)
    for d in (d1, d2):
        db.add(IssueClusterDocument(issue_cluster_id=cluster.id, document_id=d.id, membership_type=IssueMembershipType.PRIMARY, similarity_score=0.9, method=ClassificationMethod.RULE_BASED))
    db.commit()
    return cluster, d1, d2


def test_candidate_detail_has_related_documents_assignee_and_history(client, db_session):
    make_user(db_session, "ed@example.com", UserRole.RESEARCHER, "연구원A")
    other = make_user(db_session, "ed2@example.com", UserRole.RESEARCHER, "연구원B")
    make_user(db_session, "ap@example.com", UserRole.APPROVER)
    h = auth_headers(client, "ed@example.com")
    cluster, d1, d2 = _cluster_with_docs(db_session)
    created = client.post("/candidates", headers=h, json={"field_code": "ai", "title": "후보 주제 A", "source_document_id": str(d1.id),
                                                           "assignee_user_id": str(other.id)})
    assert created.status_code == 201, created.text
    cid = created.json()["id"]
    assert created.json()["issue_cluster_id"] == str(cluster.id)  # 근거 문서의 핵심 이슈가 자동 연결된다
    assert client.post("/candidates", headers=h, json={"field_code": "ai", "title": "x", "assignee_user_id": "00000000-0000-0000-0000-000000000000"}).status_code == 422
    assert client.patch(f"/candidates/{cid}", headers=h, json={"title": "후보 주제 A(수정)", "status": "hold"}).status_code == 200
    assert client.patch(f"/candidates/{cid}", headers=h, json={"status": "final"}).status_code == 403  # 확정은 승인권자만
    detail = client.get(f"/candidates/{cid}", headers=h).json()
    assert detail["assignee_name"] == "연구원B"
    docs = {d["title"]: d for d in detail["related_documents"]}
    assert set(docs) == {"Issue doc one", "Issue doc two"} and docs["Issue doc one"]["is_primary"] is True
    assert docs["Issue doc two"]["link_blocked"] is True and docs["Issue doc one"]["link_blocked"] is False
    change = next(e for e in detail["history"] if e["action"] == "candidate.update" and "title" in e["detail"]["changes"])
    assert change["detail"]["changes"]["title"] == {"from": "후보 주제 A", "to": "후보 주제 A(수정)"} and change["actor"] == "연구원A"
    assert change["detail"]["changes"]["status"]["to"] == "hold"
    # 담당자별 조회
    assert [c["title"] for c in client.get("/candidates", params={"assignee_user_id": str(other.id)}, headers=h).json()] == ["후보 주제 A(수정)"]
    assert client.get("/candidates", params={"mine": "true"}, headers=h).json() == []
    # 이슈 상세에서도 같은 후보와 원문 링크 상태를 볼 수 있다
    issue = client.get(f"/issues/{cluster.id}", headers=h).json()
    assert [c["id"] for c in issue["candidates"]] == [cid]
    assert {m["title"]: m["link_blocked"] for m in issue["members"]} == {"Issue doc one": False, "Issue doc two": True}
    assert {m["source_org"] for m in issue["members"]} == {"US Govt", "EU Commission"}


# ---------------------------------------------------------------- 요약본: 원고 핵심 사실 변경 재검토(시행일 연기 사례) ----
def _ms(client, h, sentence):
    created = client.post("/manuscripts", headers=h, json={"manuscript_type": "key_issue", "title": "시행일 점검 원고"}).json()
    return created["id"], _save(client, h, created["id"], 0, sentence)


def _save(client, h, mid, lock, text):
    doc = {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": text}]}]}
    r = client.patch(f"/manuscripts/{mid}", headers=h, json={"body_json": json.dumps(doc, ensure_ascii=False), "body": text, "lock_version": lock, "is_milestone": True})
    assert r.status_code == 200, r.text
    return r.json()["lock_version"]


def _summary(client, h, mid, sentence):
    rid = client.post("/reports", headers=h, json={"report_type": "summary", "title": "시행일 요약본", "layout": "compact"}).json()["id"]
    r = client.post(f"/reports/{rid}/summary-items", headers=h, json={"part_key": "key_issue", "manuscript_id": mid, "source_sentence": sentence})
    assert r.status_code == 201, r.text
    return rid, r.json()["items"][0]["id"]


def test_summary_keeps_old_effective_date_until_researcher_reviews_and_blocks_finalize(client, db_session):
    make_user(db_session, "w@example.com", UserRole.RESEARCHER)
    make_user(db_session, "wa@example.com", UserRole.APPROVER)
    h, ha = auth_headers(client, "w@example.com"), auth_headers(client, "wa@example.com")
    old = "최종 규정은 2025.2.18.부터 시행됨"
    mid, lock = _ms(client, h, old)
    rid, item_id = _summary(client, h, mid, old)
    assert client.get(f"/reports/{rid}", headers=h).json()["items"][0]["needs_review"] is False

    # 후속 규정으로 원고가 바뀌어도 연구원이 다듬은 요약문은 자동으로 덮어쓰이지 않는다
    new = "최종 규정의 시행일은 2025.2.18.에서 2025.3.20.으로 연기됨"
    _save(client, h, mid, lock, new)
    detail = client.get(f"/reports/{rid}", headers=h).json()
    it = detail["items"][0]
    assert it["summary_text"] == old and it["needs_review"] is True
    assert it["fact_changes"]["old_sentence"] == old and it["fact_changes"]["new_sentence"] == new
    assert "2025-03-20" in it["fact_changes"]["added"] and "2025-02-18" not in it["fact_changes"]["removed"]  # 새 사실을 알려준다
    assert any(b["code"] == "summary_stale" for b in detail["blocking"])  # 재검토 전에는 최종 확정 차단
    assert client.post(f"/reports/{rid}/finalize", headers=ha, json={}).status_code == 422
    # 사유 없이 기록할 수 없고, 기록은 연구원이 한다
    assert client.post(f"/reports/{rid}/items/{item_id}/fact-review", headers=h, json={"resolution": "no_change", "note": ""}).status_code == 422
    ok = client.post(f"/reports/{rid}/items/{item_id}/fact-review", headers=h,
                     json={"resolution": "updated", "note": "시행일 연기를 요약문에 반영함", "new_source_sentence": new})
    assert ok.status_code == 200, ok.text
    item = ok.json()["items"][0]
    assert item["needs_review"] is False and item["fact_review"]["resolution"] == "updated" and item["fact_review"]["reviewer"]
    assert not any(b["code"] == "summary_stale" for b in ok.json()["blocking"])
    # 원고가 다시 바뀌면 이전 재검토 기록으로는 통과하지 못한다
    _save(client, h, mid, client.get(f"/manuscripts/{mid}", headers=h).json()["lock_version"], new + " (추가 정정)")
    again = client.get(f"/reports/{rid}", headers=h).json()["items"][0]
    assert again["needs_review"] is False or again["fact_review"]["current"] is False


def test_past_fact_left_in_summary_cannot_be_marked_no_change(client, db_session):
    make_user(db_session, "p@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "p@example.com")
    old = "총중량 4,536kg 이하 차량의 준수 시한은 2027.9.1.임"
    mid, lock = _ms(client, h, old)
    rid, item_id = _summary(client, h, mid, old)
    new = "총중량 4,536kg 이하 차량의 준수 시한은 2028.9.1.임"
    _save(client, h, mid, lock, new)
    it = client.get(f"/reports/{rid}", headers=h).json()["items"][0]
    assert it["past_facts"] == ["2027-09-01"] and it["fact_changes"]["removed"] == ["2027-09-01"] and it["fact_changes"]["added"] == ["2028-09-01"]
    blocked = client.post(f"/reports/{rid}/items/{item_id}/fact-review", headers=h, json={"resolution": "no_change", "note": "영향 없음으로 판단"})
    assert blocked.status_code == 422 and "2027-09-01" in blocked.json()["detail"]  # 과거 사실이 남아 있으면 '영향 없음'으로 넘어갈 수 없다
    assert client.patch(f"/reports/{rid}/items/{item_id}", headers=h, json={"summary_text": new}).status_code == 200
    done = client.post(f"/reports/{rid}/items/{item_id}/fact-review", headers=h, json={"resolution": "updated", "note": "시한 정정", "new_source_sentence": new})
    assert done.status_code == 200 and done.json()["items"][0]["needs_review"] is False and done.json()["items"][0]["past_facts"] == []


def test_extract_facts_dates_and_numbers():
    from app.services.report_service import extract_facts

    assert extract_facts("2025.2.18.부터 시행, 총중량 4,536kg 이하, 38% 증가, 2024년 12월 20일 공고") == {
        "2025-02-18", "2024-12-20", "4536kg", "38%"}
    assert extract_facts("") == set() and extract_facts(None) == set()


# ---------------------------------------------------------------- 주간기술동향지: 표 제목·표 의견·심의 예외 ----
def _docx_with_caption() -> bytes:
    d = Document()
    d.add_heading("제목", level=1)
    d.add_paragraph("본문 문단입니다.")
    d.add_paragraph("[표 1] 국가별 정책 비교")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "구분", "내용", "A", "값"
    d.add_paragraph("[표 2]")
    t2 = d.add_table(rows=1, cols=1)
    t2.cell(0, 0).text = "번호만 있는 표"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_author_table_caption_paragraph_is_merged_into_table_caption_not_duplicated():
    body, skipped = journal_docx.to_editor_json(_docx_with_caption())
    types = [b["type"] for b in body["content"]]
    assert types == ["heading", "paragraph", "table", "table"]  # '[표 1] …' 문단이 표 앞에 남지 않는다
    assert body["content"][2]["attrs"] == {"caption": "국가별 정책 비교"}
    assert "attrs" not in body["content"][3]  # 번호만 있으면 제목 없이 번호만 자동 부여


def test_table_and_figure_comments_are_flagged_for_original_check_not_unchanged():
    old, new = ["문단 하나", "문단 둘"], ["문단 하나", "문단 둘"]
    mapped = journal_docx.map_comments([{"id": "1", "scope": "table", "anchor_text": "[표 1] 비교표", "text": "표 설명 보강"},
                                        {"id": "2", "scope": "paragraph", "anchor_text": "문단 하나", "text": "근거 보강"}], old, new)
    assert mapped[0]["match"] == "needs_original_check" and "원본 파일" in mapped[0]["note"]
    assert mapped[1]["match"] == "unchanged"


def test_publish_with_unresolved_required_comments_needs_approver_exception_reason(client, db_session):
    from tests.test_journal import World, docx_file

    w = World(client, db_session)
    e, author_h = w.expert_with_account()
    iid = w.issue()
    com = w.accepted_commission(iid, e)
    ms = w.submit(author_h, com).json()
    mid = ms["id"]
    client.post(f"/journal/manuscripts/{mid}/format-review", headers=w.ed, json={"result": "ok"})
    rev = client.get("/journal/reviewer-candidates", headers=w.ed).json()
    reviewer = next(r for r in rev if r["display_name"] == "내부검토자")
    a = client.post(f"/journal/manuscripts/{mid}/assignments", headers=w.ed, json={"reviewer_user_id": reviewer["id"]}).json()
    client.post(f"/journal/portal/reviews/{a['id']}/conflict", headers=w.rv, json={"declared": False})
    client.post(f"/journal/portal/reviews/{a['id']}/comments", headers=w.rv, json={"scope": "table", "anchor_text": "[표 1]", "text": "표 정의를 보강해 주세요"})
    client.post(f"/journal/portal/reviews/{a['id']}/complete", headers=w.rv, json={"recommendation": "accept", "summary_opinion": "적합"})
    client.post(f"/journal/manuscripts/{mid}/deliberate", headers=w.ed)
    # 저자는 답변만 했다 → 해결이 아니므로 확정이 막힌다
    cm = client.get(f"/journal/portal/author/manuscripts/{mid}", headers=author_h).json()["comments"]
    assert cm and cm[0]["status"] == "open"
    blocked = client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "accept"})
    assert blocked.status_code == 409 and "승인권자 예외 사유" in blocked.json()["detail"]
    dash = client.get("/journal/dashboard?year=2027", headers=w.ed).json()
    assert dash["unresolved_comments"][0]["count"] == 1  # 미해결 의견이 대시보드에 보인다
    # 예외 승인: 승인권자와 사유가 기록되고 이후 대시보드에서 예외 확정을 확인할 수 있다
    ok = client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "accept", "exception_reason": "표 의견은 교정 단계에서 처리"})
    assert ok.status_code == 200, ok.text
    last = ok.json()["decisions"][-1]
    assert last["exception_reason"] == "표 의견은 교정 단계에서 처리" and last["unresolved_count"] == 1 and last["decided_by"] == "승인권자"
    assert client.get("/journal/dashboard?year=2027", headers=w.ed).json()["issues"][0]["exception_decisions"] == 1


# ---------------------------------------------------------------- 공개 스냅샷 개인정보·긴 제목 줄바꿈 ----
def test_public_snapshot_excludes_personal_data_text_and_test_items(db_session):
    from app import export_public_snapshot as snap

    make_doc(db_session, title="Real public AI policy contact a.b@corp.com", field=FieldCode.AI, url="https://example.org/real", full_text="원문 전문 secret-body")
    make_doc(db_session, title="[시범] 테스트 문서", field=FieldCode.AI)
    make_doc(db_session, title="E2E 문서", field=FieldCode.AI)
    make_doc(db_session, title="Skip me", field=FieldCode.AI, relevance=ClassificationRelevance.LOW_RELEVANCE)
    s = snap.build_snapshot(db_session)
    titles = [d["title"] for d in s["documents"]]
    assert len(titles) == 1 and "@" not in titles[0]  # 시범·E2E·무관 문서 제외, 이메일 제거
    text = json.dumps(s, ensure_ascii=False)
    assert "secret-body" not in text  # 원문 전문은 포함하지 않는다
    assert s["documents"][0]["original_secured"] is True  # 확보 여부만 표시
    assert {"summary", "full_text", "reviews", "email"}.isdisjoint(s["documents"][0].keys())
    with pytest.raises(snap.PrivacyError):
        snap.assert_no_personal_data({"x": "문의 someone@example.com"})


def test_long_title_wraps_at_word_boundaries():
    from app.services.report_docx import _wrap_by_words

    out = _wrap_by_words("[시범] 주간 AI·ICT 동향 요약본(간결형) — 개발용 출력 테스트", 22, 8.0)
    lines = out.split("\n")
    assert len(lines) >= 2 and "".join(lines).replace(" ", "") == "[시범]주간AI·ICT동향요약본(간결형)—개발용출력테스트"
    assert not any(l.startswith(" ") for l in lines)


# ---------------------------------------------------------------- Phase 3-7: 후보주제 상세(이슈 요약·논쟁 가능성)·근거 추가/제외·분야 수정·추가 조사 필요 ----
def test_candidate_detail_issue_summary_dispute_label_evidence_add_exclude_field_and_research_status(client, db_session):
    make_user(db_session, "c37@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "c37@example.com")
    cluster, d1, d2 = _cluster_with_docs(db_session)
    extra = make_doc(db_session, title="Extra evidence doc", field=FieldCode.AI, url="https://example.org/extra", days_ago=1)
    cid = client.post("/candidates", headers=h, json={"field_code": "ai", "title": "후보 37", "source_document_id": str(d1.id)}).json()["id"]
    detail = client.get(f"/candidates/{cid}", headers=h).json()
    s = detail["issue_summary"]
    assert s["issue_id"] == str(cluster.id) and s["document_count"] == 2 and s["dispute_status"] == "not_analyzed" and "논쟁으로 표시하지 않습니다" in s["dispute_note"]
    assert detail["domestic_relevance_note"] and "평가 보류" in detail["domestic_relevance_note"]
    assert {d["title"] for d in detail["related_documents"]} == {"Issue doc one", "Issue doc two"}
    # 근거 원문 제외·복원(대표 근거는 제외 불가)·추가
    assert client.post(f"/candidates/{cid}/evidence", headers=h, json={"document_id": str(d1.id), "action": "exclude"}).status_code == 409
    ex = client.post(f"/candidates/{cid}/evidence", headers=h, json={"document_id": str(d2.id), "action": "exclude", "note": "관련성 낮음"}).json()
    assert [d["title"] for d in ex["related_documents"]] == ["Issue doc one"] and [d["title"] for d in ex["excluded_documents"]] == ["Issue doc two"]
    assert client.post(f"/candidates/{cid}/evidence", headers=h, json={"document_id": str(d2.id), "action": "restore"}).json()["excluded_documents"] == []
    add = client.post(f"/candidates/{cid}/evidence", headers=h, json={"document_id": str(extra.id), "action": "add"}).json()
    assert any(d["title"] == "Extra evidence doc" and d["added_by_researcher"] for d in add["related_documents"])
    assert client.post(f"/candidates/{cid}/evidence", headers=h, json={"document_id": str(extra.id), "action": "add"}).status_code == 409
    # 분야 수정·추가 조사 필요 상태·이력
    assert client.patch(f"/candidates/{cid}", headers=h, json={"field_code": "sec", "status": "research", "title": "후보 37(수정)"}).status_code == 200
    after = client.get(f"/candidates/{cid}", headers=h).json()
    assert after["field_code"] == "sec" and after["status"] == "research" and after["title"] == "후보 37(수정)"
    actions = [e["action"] for e in after["history"]]
    assert "candidate.update" in actions and "candidate.evidence.exclude" in actions and "candidate.evidence.add" in actions


def test_manuscript_lint_checks_values_roles_duplicates_and_never_changes_the_manuscript(client, db_session):
    make_user(db_session, "lint37@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "lint37@example.com")
    ms = client.post("/manuscripts", headers=h, json={"manuscript_type": "key_issue", "title": "점검 원고"}).json()
    body_text = "최종 규정은 2025.3.20.부터 시행됨. 준수 시한은 2025.2.18.임. 같은 문장이 반복됨. 같은 문장이 반복됨. 비용은 120억 원으로 추산됨."
    doc = {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": body_text}]}]}
    saved = client.patch(f"/manuscripts/{ms['id']}", headers=h, json={"body_json": json.dumps(doc, ensure_ascii=False), "body": body_text, "lock_version": ms["lock_version"]}).json()
    q = "The final rule takes effect on February 18, 2025. Manufacturers must comply by December 22, 2025."
    for anchor, mark in (("최종 규정은 2025.3.20.부터 시행됨.", "m1"), ("준수 시한은 2025.2.18.임.", "m2")):
        assert client.post(f"/manuscripts/{ms['id']}/footnotes", headers=h, json={"mark_id": mark, "anchor_text": anchor, "source_url": "https://example.org/fr", "excerpt": q}).status_code == 201
    assert client.post(f"/manuscripts/{ms['id']}/footnotes", headers=h, json={"mark_id": "m3", "anchor_text": "발췌 없는 문장입니다.", "source_url": "https://example.org/fr"}).status_code == 201
    out = client.get(f"/manuscripts/{ms['id']}/lint", headers=h).json()
    kinds = [(w["type"], w["severity"]) for w in out["warnings"]]
    assert ("value_mismatch", "error") in kinds  # 2025.3.20.은 발췌문에 없음
    assert any(w["type"] == "semantic_risk" and "역할이 다름" in w["message"] for w in out["warnings"])  # 시행일 문장 vs 준수 시한 혼동
    assert any(w["type"] == "duplicate_sentence" for w in out["warnings"]) and any(w["type"] == "duplicate_source" for w in out["warnings"])
    assert any(w["type"] == "numeric_without_source" and "120억" in w["sentence"] for w in out["warnings"])
    assert any(w["type"] == "no_excerpt" for w in out["warnings"]) and any(w["type"] == "unverified" for w in out["warnings"])
    assert out["counts"]["error"] >= 1 and "자동으로 고치거나" in out["note"]
    after = client.get(f"/manuscripts/{ms['id']}", headers=h).json()
    assert after["body"] == body_text and after["lock_version"] == saved["lock_version"]  # 점검은 원고를 바꾸지 않는다
    assert all(f["verification_status"] == "not_verified" for f in client.get(f"/manuscripts/{ms['id']}/footnotes", headers=h).json())


def test_dashboard_counts_reconcile(client, db_session):
    """전체 자료 = 관련 + 관련성 낮음 + 분류 보류 + 미처리, 관련 = 분야별 합계 + 주 분야 미지정."""
    from tests.conftest import auth_headers, make_user
    from app.models.user import UserRole
    make_user(db_session, "recon-admin@example.com", UserRole.ADMIN)
    s = client.get("/dashboard/summary", headers=auth_headers(client, "recon-admin@example.com")).json()
    assert s["total_documents"] == s["matched_documents"] + s["low_relevance_documents"] + s["pending_review_documents"] + s["unclassified_documents"]
    assert s["matched_documents"] == sum(f["document_count"] for f in s["by_field"]) + s["matched_without_field"]
    assert s["issue_cluster_without_candidate"] >= 0


def test_priority_candidates_exclude_admin_notices_and_cap(client, db_session):
    from app.services.candidate_priority import admin_notice_reason, MAX_ITEMS
    assert admin_notice_reason("Radio Broadcasting Services; Selmer, Tennessee")
    assert admin_notice_reason("Petition for Reconsideration of Action in Rulemaking Proceeding")
    assert admin_notice_reason("Upper C-Band (3.98-4.2 GHz); Expanding Flexible Use") is None
    make_user(db_session, "prio-admin@example.com", UserRole.ADMIN)
    r = client.get("/candidates/priority", headers=auth_headers(client, "prio-admin@example.com"))
    assert r.status_code == 200
    body = r.json()
    assert len(body["items"]) <= MAX_ITEMS and body["criteria"]
