"""Phase 3-3 — 카드뉴스 생성·편집·근거·버전·검토·PPTX/PNG 출력·권한 테스트.

생성된 PPTX는 python-pptx로 다시 열어 슬라이드 크기·텍스트 객체·표·차트(임베디드 데이터)를 확인하고, PNG는 LibreOffice로 실제
렌더링한 결과의 해상도·비율을 확인한다(렌더링 테스트는 LibreOffice가 있을 때만, CI에서는 필수).
"""

import io
import json
import os
import zipfile

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Cm

from app.core.config import settings
from app.models.enums import UserRole
from app.models.card_news import CardNewsExport
from app.services import card_news_pptx, card_news_scene, report_pdf
from app.services.card_news_schema import template_config
from tests.conftest import auth_headers, make_user
from tests.test_reports import _general_report, _make_manuscript

_has_soffice = report_pdf.find_soffice() is not None
requires_soffice = pytest.mark.skipif(not _has_soffice, reason="LibreOffice 없음 — CI에서 검증")


@pytest.fixture(autouse=True)
def _export_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "export_dir", str(tmp_path / "exports"))


@pytest.fixture(autouse=True)
def _skip_rendering_unless_marked(request, monkeypatch):
    if request.node.get_closest_marker("render") is None:
        monkeypatch.setattr(report_pdf, "find_soffice", lambda: None)


@pytest.mark.render
def test_libreoffice_available_when_required_by_ci():
    if os.environ.get("REQUIRE_SOFFICE") == "1":
        assert report_pdf.find_soffice() is not None


def _create(client, h, template="A", title="테스트 카드뉴스", **extra):
    resp = client.post("/card-news", headers=h, json={"title": title, "template_key": template, **extra})
    assert resp.status_code == 201, resp.text
    return resp.json()


def _set(client, h, cn, card_index, key, **body):
    card = cn["cards"][card_index]
    resp = client.patch(f"/card-news/{cn['id']}/cards/{card['id']}/elements/{key}", headers=h, json={"lock_version": cn["lock_version"], **body})
    assert resp.status_code == 200, resp.text
    return resp.json()


def _el(cn, card_index, key):
    return next(e for e in cn["cards"][card_index]["elements"] if e["key"] == key)


def _fill_A(client, h, cn):
    cn = _set(client, h, cn, 1, "heading", text="추진 경과")
    cn = _set(client, h, cn, 1, "content_1", text="미국 NHTSA는 2024.4.15. 제정안을 공고함")
    cn = _set(client, h, cn, 1, "source", text="NHTSA")
    cn = _set(client, h, cn, 0, "title", text="전기차 안전기준 동향")
    cn = _set(client, h, cn, 2, "message", text="일정 구분 관리가 필요함")
    cn = _set(client, h, cn, 2, "source", text="NHTSA")
    return cn


# ---------------------------------------------------------------- 템플릿·기본 ----
def test_endpoints_require_auth(client):
    assert client.get("/card-news").status_code == 401
    assert client.post("/card-news", json={"title": "x", "template_key": "A"}).status_code == 401


def test_four_templates_have_different_structures(client, db_session):
    make_user(db_session, "t1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "t1@example.com")
    templates = {t["key"]: t for t in client.get("/card-news/templates", headers=h).json()}
    assert set(templates) == {"A", "B", "C", "D"}
    body_keys = {k: {e["key"] for e in t["elements"]["body"]} for k, t in templates.items()}
    assert "content_3" in body_keys["A"] and "stage_5_date" in body_keys["B"]
    assert "country_3_policy" in body_keys["C"] and "chart" in body_keys["D"]
    assert len({frozenset(v) for v in body_keys.values()}) == 4  # 색만 다른 것이 아니라 요소 구성이 모두 다르다
    assert templates["A"]["config"]["width_cm"] == 7.0 and templates["A"]["config"]["height_cm"] == 15.62
    assert templates["A"]["config"]["navy"] == "1B2A4A"


@pytest.mark.parametrize("key", ["A", "B", "C", "D"])
def test_manual_creation_makes_cover_body_closing_and_flags_empty_required(client, db_session, key):
    make_user(db_session, f"m{key}@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, f"m{key}@example.com")
    cn = _create(client, h, key, body_count=2)
    assert [c["kind"] for c in cn["cards"]] == ["cover", "body", "body", "closing"]
    assert cn["width_cm"] == 7.0 and cn["height_cm"] == 15.62 and cn["status"] == "draft"
    assert any(i["code"] == "empty_required" for i in cn["preflight"])
    assert {b["code"] for b in cn["blocking"]} >= {"empty_required"}
    assert "공식 발간 규격" in cn["size_note"]
    scenes = client.get(f"/card-news/{cn['id']}/scenes", headers=h).json()["scenes"]
    assert len(scenes) == 4 and all(s["width"] == 7.0 for s in scenes)


# ---------------------------------------------------------------- 보고서 연계 ----
def _report_with_manuscript(client, h, *, verify=True):
    mid = _make_manuscript(client, h, title="FMVSS 시범 원고", verify=verify)
    rid = _general_report(client, h, {"key_issue": mid})
    return mid, rid


def test_create_from_report_uses_only_real_sentences_and_keeps_evidence_unverified(client, db_session):
    make_user(db_session, "r1@example.com", UserRole.RESEARCHER)
    make_user(db_session, "r2@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r1@example.com")
    mid, rid = _report_with_manuscript(client, h, verify=True)  # 원고 각주는 검증 완료 상태

    cand = client.get(f"/card-news/report-sentences/{rid}", headers=h).json()
    target = next(s for s in cand["sentences"] if "제정안을 공고함" in s["sentence"])
    assert target["footnotes"] and target["footnotes"][0]["status"] == "matched"

    created = client.post("/card-news", headers=h, json={
        "title": "보고서 연계 카드뉴스", "template_key": "A", "source_type": "report", "source_report_id": rid,
        "selections": [{"manuscript_id": mid, "sentence": target["sentence"]}],
    })
    assert created.status_code == 201, created.text
    cn = created.json()
    body_el = _el(cn, 1, "content_1")
    assert body_el["text"] == target["sentence"] and body_el["source_manuscript_id"] == mid
    assert body_el["evidence"] and body_el["evidence"][0]["original_footnote_status"] == "matched"
    # 원고 각주가 검증되었어도 카드에서는 연구원이 따로 확인하기 전까지 검증 완료가 아니다
    assert body_el["evidence_state"] == "not_verified"
    assert any(i["code"] == "unverified_claim" for i in cn["preflight"])
    assert cn["is_pilot"] is False and cn["source_changed"] is False

    forged = client.post("/card-news", headers=h, json={
        "title": "x", "template_key": "A", "source_type": "report", "source_report_id": rid,
        "selections": [{"manuscript_id": mid, "sentence": "원고에 없는 지어낸 문장입니다"}]})
    assert forged.status_code == 422
    h2 = auth_headers(client, "r2@example.com")
    assert client.get(f"/card-news/report-sentences/{rid}", headers=h2).status_code == 403
    assert client.get(f"/card-news/{cn['id']}", headers=h2).status_code == 403


def test_timeline_template_extracts_only_dates_present_in_sentence(client, db_session):
    make_user(db_session, "r3@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r3@example.com")
    mid, rid = _report_with_manuscript(client, h)
    sents = client.get(f"/card-news/report-sentences/{rid}", headers=h).json()["sentences"]
    dated = next(s for s in sents if "2025.2.18." in s["sentence"])
    undated = next(s for s in sents if "2025.2.18." not in s["sentence"] and "굵은" in s["sentence"])
    cn = client.post("/card-news", headers=h, json={
        "title": "일정", "template_key": "B", "source_type": "report", "source_report_id": rid,
        "selections": [{"manuscript_id": mid, "sentence": dated["sentence"]}, {"manuscript_id": mid, "sentence": undated["sentence"]}]}).json()
    assert _el(cn, 1, "stage_1_date")["text"] == "2025.2.18."
    assert _el(cn, 1, "stage_2_date")["text"] == ""  # 날짜가 없는 문장에서 날짜를 만들어 넣지 않는다


def test_source_change_marks_review_needed_without_overwriting_edited_card(client, db_session):
    make_user(db_session, "r4@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "r4@example.com")
    mid, rid = _report_with_manuscript(client, h)
    sent = next(s for s in client.get(f"/card-news/report-sentences/{rid}", headers=h).json()["sentences"] if "제정안을 공고함" in s["sentence"])
    cn = client.post("/card-news", headers=h, json={"title": "변경 감지", "template_key": "A", "source_type": "report",
                                                      "source_report_id": rid, "selections": [{"manuscript_id": mid, "sentence": sent["sentence"]}]}).json()
    cn = _set(client, h, cn, 1, "content_1", text="연구원이 다듬은 핵심 내용")
    ms = client.get(f"/manuscripts/{mid}", headers=h).json()
    changed = {"type": "doc", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "완전히 다른 내용으로 원고가 바뀌었습니다"}]}]}
    client.patch(f"/manuscripts/{mid}", headers=h, json={"body_json": json.dumps(changed, ensure_ascii=False),
                                                         "body": "완전히 다른 내용으로 원고가 바뀌었습니다", "lock_version": ms["lock_version"]})
    after = client.get(f"/card-news/{cn['id']}", headers=h).json()
    assert after["needs_review"] is True and after["source_changed"] is True
    assert _el(after, 1, "content_1")["text"] == "연구원이 다듬은 핵심 내용"  # 자동으로 덮어쓰지 않음
    assert _el(after, 1, "content_1")["source_stale"] is True
    assert any(i["code"] == "source_stale" for i in after["preflight"])


# ---------------------------------------------------------------- 편집·충돌·근거 ----
def test_edit_persists_and_conflict_detected(client, db_session):
    make_user(db_session, "e1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "e1@example.com")
    cn = _create(client, h, "A")
    first = _set(client, h, cn, 1, "heading", text="저장 확인")
    assert first["lock_version"] == cn["lock_version"] + 1
    reloaded = client.get(f"/card-news/{cn['id']}", headers=h).json()
    assert _el(reloaded, 1, "heading")["text"] == "저장 확인"
    stale = client.patch(f"/card-news/{cn['id']}/cards/{cn['cards'][1]['id']}/elements/heading", headers=h,
                         json={"lock_version": cn["lock_version"], "text": "오래된 수정"})
    assert stale.status_code == 409 and "다른 사용자" in stale.json()["detail"]
    style = _set(client, h, first, 1, "heading", style={"size": 12, "color": "FFFFFF", "bold": True, "align": "center"})
    assert _el(style, 1, "heading")["style"]["size"] == 12


def test_evidence_verification_requires_explicit_action_and_goes_stale_after_edit(client, db_session):
    make_user(db_session, "e2@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "e2@example.com")
    cn = _create(client, h, "A")
    cn = _set(client, h, cn, 1, "content_1", text="2025.2.18.부터 시행됨")
    card_id = cn["cards"][1]["id"]
    added = client.post(f"/card-news/{cn['id']}/cards/{card_id}/evidence", headers=h, json={
        "element_key": "content_1", "source_url": "https://example.org/rule", "source_org": "NHTSA", "excerpt": "effective February 18, 2025"})
    assert added.status_code == 201
    el = _el(added.json(), 1, "content_1")
    assert el["evidence_state"] == "not_verified"
    ev_id = el["evidence"][0]["id"]
    verified = client.post(f"/card-news/{cn['id']}/evidence/{ev_id}/verify", headers=h, json={"verified": True}).json()
    assert _el(verified, 1, "content_1")["evidence_state"] == "verified"
    edited = _set(client, h, verified, 1, "content_1", text="2025.3.18.부터 시행됨")  # 날짜를 고치면 검증이 무효가 된다
    assert _el(edited, 1, "content_1")["evidence_state"] == "stale"
    assert any(i["code"] == "evidence_stale" for i in edited["preflight"])
    bad = client.post(f"/card-news/{cn['id']}/cards/{card_id}/evidence", headers=h, json={"element_key": "source", "source_url": "https://x.org"})
    assert bad.status_code == 422  # 근거 대상이 아닌 요소


def test_card_add_duplicate_reorder_delete_and_template_change(client, db_session):
    make_user(db_session, "e3@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "e3@example.com")
    cn = _create(client, h, "A", body_count=1)
    cn = _set(client, h, cn, 1, "content_1", text="이동되는 핵심 내용")
    cn = _set(client, h, cn, 1, "heading", text="제목 유지")
    assert len(cn["cards"]) == 3
    cn = client.post(f"/card-news/{cn['id']}/cards", headers=h, json={"lock_version": cn["lock_version"], "kind": "body", "after_card_id": cn["cards"][1]["id"]}).json()
    assert len(cn["cards"]) == 4 and cn["cards"][2]["kind"] == "body"
    cn = client.post(f"/card-news/{cn['id']}/cards/{cn['cards'][1]['id']}/duplicate", headers=h, json={"lock_version": cn["lock_version"]}).json()
    assert len(cn["cards"]) == 5 and _el(cn, 2, "content_1")["text"] == "이동되는 핵심 내용"
    ids = [c["id"] for c in cn["cards"]]
    cn = client.post(f"/card-news/{cn['id']}/cards/reorder", headers=h, json={"lock_version": cn["lock_version"], "card_ids": list(reversed(ids))}).json()
    assert [c["id"] for c in cn["cards"]] == list(reversed(ids))
    cn = client.delete(f"/card-news/{cn['id']}/cards/{cn['cards'][1]['id']}", headers=h, params={"lock_version": cn["lock_version"]}).json()
    assert len(cn["cards"]) == 4
    changed = client.patch(f"/card-news/{cn['id']}", headers=h, json={"lock_version": cn["lock_version"], "template_key": "B"}).json()
    assert changed["template_key"] == "B"
    moved = [e for c in changed["cards"] for e in c["elements"] if e["key"] == "stage_1_text" and e["text"]]
    assert moved and moved[0]["text"] == "이동되는 핵심 내용"  # 템플릿을 바꿔도 내용이 보존된다
    assert any(e["key"] == "stage_1_kind" for c in changed["cards"] if c["kind"] == "body" for e in c["elements"])
    single = _create(client, h, "A", body_count=1)
    for _ in range(2):  # 3장 → 1장
        single = client.delete(f"/card-news/{single['id']}/cards/{single['cards'][-1]['id']}", headers=h,
                               params={"lock_version": single["lock_version"]}).json()
    assert len(single["cards"]) == 1
    last = client.delete(f"/card-news/{single['id']}/cards/{single['cards'][0]['id']}", headers=h, params={"lock_version": single["lock_version"]})
    assert last.status_code == 422


def test_versions_milestone_and_restore(client, db_session):
    make_user(db_session, "v1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "v1@example.com")
    cn = _create(client, h, "A")
    cn = _set(client, h, cn, 1, "heading", text="원래 제목")
    saved = client.post(f"/card-news/{cn['id']}/versions", headers=h, json={"change_note": "1차 완료"})
    assert saved.status_code == 201
    cn = _set(client, h, cn, 1, "heading", text="바뀐 제목")
    versions = client.get(f"/card-news/{cn['id']}/versions", headers=h).json()
    milestone = next(v for v in versions if v["change_note"] == "1차 완료")
    restored = client.post(f"/card-news/{cn['id']}/versions/{milestone['id']}/restore", headers=h, json={"lock_version": cn["lock_version"]})
    assert restored.status_code == 200, restored.text
    assert _el(restored.json(), 1, "heading")["text"] == "원래 제목"
    after = client.get(f"/card-news/{cn['id']}/versions", headers=h).json()
    assert any("복원 전 상태" in (v["change_note"] or "") for v in after)  # 복원 전 상태도 되돌릴 수 있게 남는다


# ---------------------------------------------------------------- 검토·승인·권한 ----
def test_review_approval_workflow_blocks_incomplete_and_separates_iitp(client, db_session):
    author = make_user(db_session, "w1@example.com", UserRole.RESEARCHER, "작성자")
    make_user(db_session, "w2@example.com", UserRole.REVIEWER, "검토자")
    make_user(db_session, "w3@example.com", UserRole.APPROVER, "승인권자")
    ha, hr, hb = (auth_headers(client, e) for e in ("w1@example.com", "w2@example.com", "w3@example.com"))
    cn = _fill_A(client, ha, _create(client, ha, "A", body_count=1))
    reviewer = next(u for u in client.get("/users", headers=ha).json() if u["role"] == "reviewer")
    assert client.post(f"/card-news/{cn['id']}/approve", headers=hb, json={}).status_code == 409  # 검토 요청 전
    rv = client.post(f"/card-news/{cn['id']}/request-review", headers=ha, json={"reviewer_user_id": reviewer["id"]}).json()
    assert rv["status"] == "review_requested"
    assert client.post(f"/card-news/{cn['id']}/request-changes", headers=ha, json={"comment_text": "x"}).status_code == 403
    chg = client.post(f"/card-news/{cn['id']}/request-changes", headers=hr, json={"comment_text": "수치 근거를 보강해주세요"}).json()
    assert chg["status"] == "changes_requested"
    reviews = client.get(f"/card-news/{cn['id']}/reviews", headers=ha).json()
    assert reviews[0]["status"] == "open"
    client.post(f"/card-news/{cn['id']}/reviews/{reviews[0]['id']}/resolve", headers=ha)
    client.post(f"/card-news/{cn['id']}/request-review", headers=ha, json={})
    assert client.post(f"/card-news/{cn['id']}/approve", headers=ha, json={}).status_code == 403  # 연구원은 확정 불가
    blocked = client.post(f"/card-news/{cn['id']}/approve", headers=hb, json={})
    assert blocked.status_code == 422  # 미검증 근거·미완성 요소가 남아 있음
    assert blocked.json()["detail"]["blocking"]
    ok = client.post(f"/card-news/{cn['id']}/approve", headers=hb, json={"exception_reason": "내부 점검용 예외"})
    assert ok.status_code == 200 and ok.json()["status"] == "approved"
    assert ok.json()["iitp_confirmed"] is False  # 내부 확정 ≠ IITP 확인
    # 확정 후 수정 차단 → 승인 취소 또는 새 수정 버전
    cur = client.get(f"/card-news/{cn['id']}", headers=ha).json()
    locked = client.patch(f"/card-news/{cn['id']}/cards/{cur['cards'][1]['id']}/elements/heading", headers=ha, json={"lock_version": cur["lock_version"], "text": "x"})
    assert locked.status_code == 409
    assert client.post(f"/card-news/{cn['id']}/reopen", headers=ha).status_code == 403
    copy = client.post(f"/card-news/{cn['id']}/duplicate", headers=ha)
    assert copy.status_code == 201 and copy.json()["status"] == "draft" and "수정본" in copy.json()["title"]
    iitp = client.patch(f"/card-news/{cn['id']}/iitp", headers=hb, json={"confirmed": True, "note": "IITP 담당자 확인"}).json()
    assert iitp["iitp_confirmed"] is True and iitp["status"] == "approved"
    assert client.patch(f"/card-news/{cn['id']}/iitp", headers=ha, json={"confirmed": False}).status_code == 403
    reopened = client.post(f"/card-news/{cn['id']}/reopen", headers=hb).json()
    assert reopened["status"] == "draft"


# ---------------------------------------------------------------- 장면·PPTX ----
def _scene_for(template, kind, elements, **ctx):
    cfg = template_config(template)
    card = {"kind": kind, "bg_color": None, "elements": elements, "index": 2, "total": 4}
    base = {"for_export": True, "labels": {}, "is_pilot": False, "is_synthetic": False, "iitp_confirmed": False, "chart_verified": False}
    base.update(ctx)
    return card_news_scene.build_card_scene(cfg["width_cm"], cfg["height_cm"], cfg, template, card, base), cfg


def _t(text):
    return {"type": "text", "text": text, "data": None, "style": {}, "box": {}, "hidden": False}


def test_scene_check_flags_overflow_tiny_font_and_low_contrast():
    long_text = "아주 긴 내용 " * 40
    scene, cfg = _scene_for("A", "body", {"heading": _t("제목"), "content_1": _t(long_text), "source": _t("출처")})
    codes = {i["code"] for i in card_news_scene.check_scene(scene, cfg)}
    assert "text_overflow" in codes
    tiny, cfg = _scene_for("A", "body", {"heading": _t("제목"), "content_1": {**_t("짧음"), "style": {"size": 5}}, "source": _t("출처")})
    assert "tiny_font" in {i["code"] for i in card_news_scene.check_scene(tiny, cfg)}
    low, cfg = _scene_for("A", "body", {"heading": _t("제목"), "content_1": {**_t("연한 글자"), "style": {"color": "F0F0F0"}}, "source": _t("출처")})
    assert "low_contrast" in {i["code"] for i in card_news_scene.check_scene(low, cfg)}
    ok, cfg = _scene_for("A", "body", {"heading": _t("제목"), "content_1": _t("정상 길이의 핵심 내용"), "source": _t("출처")})
    assert not [i for i in card_news_scene.check_scene(ok, cfg) if i["level"] == "error"]


def test_chart_only_when_verified_or_synthetic_otherwise_message_layout():
    data = {"title": "비중", "unit": "%", "base_year": "2026", "comparison": "국가별", "source": "테스트", "synthetic": False,
            "series": [{"label": "A", "value": 10}, {"label": "B", "value": 20}]}
    chart_el = {"type": "chart", "text": None, "data": data, "style": {}, "box": {}, "hidden": False}
    els = {"heading": _t("수치"), "stat_value": _t("20"), "meaning": _t("의미"), "message": _t("메시지"), "source": _t("출처"), "chart": chart_el}
    unverified, _ = _scene_for("D", "body", els, chart_verified=False)
    assert unverified["layout"] == "message" and not any(i["type"] == "chart" for i in unverified["items"])
    verified, _ = _scene_for("D", "body", els, chart_verified=True)
    assert verified["layout"] == "chart" and any(i["type"] == "chart" for i in verified["items"])
    synth, _ = _scene_for("D", "body", {**els, "chart": {**chart_el, "data": {**data, "synthetic": True}}}, chart_verified=False)
    assert synth["layout"] == "chart"
    incomplete, _ = _scene_for("D", "body", {**els, "chart": {**chart_el, "data": {**data, "unit": "", "synthetic": True}}})
    assert incomplete["layout"] == "message"  # 단위가 없으면 합성 데이터여도 그래프를 그리지 않는다
    assert card_news_scene.validate_chart({**data, "base_year": ""}) == ["기준연도가 없습니다"]


def _open_pptx(data):
    return Presentation(io.BytesIO(data))


def test_pptx_export_is_real_editable_objects_not_images(client, db_session):
    make_user(db_session, "p1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "p1@example.com")
    cn = _fill_A(client, h, _create(client, h, "A", body_count=1))
    exp = client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "pptx"})
    assert exp.status_code == 201, exp.text
    e = exp.json()
    assert e["status"] == "success" and e["card_count"] == 3 and e["file_size"] > 5000
    dl = client.get(f"/card-news/{cn['id']}/exports/{e['id']}/download", headers=h)
    import hashlib

    assert hashlib.sha256(dl.content).hexdigest() == e["file_sha256"]
    prs = _open_pptx(dl.content)
    assert len(prs.slides) == 3
    assert abs(prs.slide_width - Cm(7.0)) < 400 and abs(prs.slide_height - Cm(15.62)) < 400  # 슬라이드 크기가 규격과 정확히 일치
    for slide in prs.slides:
        assert not [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]  # 전체 이미지 삽입 방식이 아님
        assert len([s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]) >= 2
    texts = [s.text_frame.text for sl in prs.slides for s in sl.shapes if s.has_text_frame]
    assert any("전기차 안전기준 동향" in t for t in texts) and any("2024.4.15." in t for t in texts)
    names = {s.name for sl in prs.slides for s in sl.shapes}
    assert {"title", "content_1", "heading"} <= names  # 요소 ID가 도형 이름으로 남는다
    runs = [r for sl in prs.slides for s in sl.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]
    assert all(r.font.name == "맑은 고딕" for r in runs) and all(r.font.size is not None for r in runs)


def test_pptx_has_editable_table_and_chart_with_embedded_data(client, db_session):
    make_user(db_session, "p2@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "p2@example.com")
    c = _create(client, h, "C", title="[합성 테스트 데이터] 비교", body_count=1, source_type="manual")
    c = _set(client, h, c, 1, "heading", text="국가 비교")
    c = _set(client, h, c, 1, "country_1_name", text="A국")
    c = _set(client, h, c, 1, "country_1_policy", text="합성 정책 설명")
    table_pptx = client.post(f"/card-news/{c['id']}/exports", headers=h, json={"file_format": "pptx"}).json()
    prs = _open_pptx(client.get(f"/card-news/{c['id']}/exports/{table_pptx['id']}/download", headers=h).content)
    tables = [s for sl in prs.slides for s in sl.shapes if s.has_table]
    assert tables and tables[0].table.cell(1, 0).text == "A국" and tables[0].table.cell(1, 1).text == "합성 정책 설명"

    d = _create(client, h, "D", title="합성 수치", body_count=1)
    chart_data = {"title": "합성 비중", "unit": "%", "base_year": "2026", "comparison": "국가별", "source": "합성 테스트 데이터", "synthetic": True,
                  "series": [{"label": "A국", "value": 38}, {"label": "B국", "value": 24}]}
    d = _set(client, h, d, 1, "chart", data=chart_data)
    d = _set(client, h, d, 1, "heading", text="수치")
    chart_pptx = client.post(f"/card-news/{d['id']}/exports", headers=h, json={"file_format": "pptx"}).json()
    prs2 = _open_pptx(client.get(f"/card-news/{d['id']}/exports/{chart_pptx['id']}/download", headers=h).content)
    charts = [s for sl in prs2.slides for s in sl.shapes if s.has_chart]
    assert charts, "막대 그래프가 PowerPoint 차트 객체로 있어야 한다"
    chart = charts[0].chart
    assert list(chart.plots[0].categories) == ["A국", "B국"] and list(chart.plots[0].series[0].values) == [38.0, 24.0]
    assert chart.part.chart_workbook.xlsx_part is not None  # 데이터 편집 가능(임베디드 엑셀)


def test_export_failure_for_png_without_libreoffice_is_recorded_and_other_user_blocked(client, db_session):
    make_user(db_session, "p3@example.com", UserRole.RESEARCHER)
    make_user(db_session, "p4@example.com", UserRole.RESEARCHER)
    h, hx = auth_headers(client, "p3@example.com"), auth_headers(client, "p4@example.com")
    cn = _fill_A(client, h, _create(client, h, "A", body_count=1))
    failed = client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "png"}).json()
    assert failed["status"] == "failed" and "LibreOffice" in failed["error_message"]
    assert client.get(f"/card-news/{cn['id']}/exports/{failed['id']}/download", headers=h).status_code == 404
    ok = client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "pptx"}).json()
    assert ok["status"] == "success"
    assert client.get(f"/card-news/{cn['id']}/exports/{ok['id']}/download", headers=hx).status_code == 403
    assert client.get(f"/card-news/{cn['id']}/exports", headers=hx).status_code == 403
    assert client.patch(f"/card-news/{cn['id']}/cards/{cn['cards'][0]['id']}/elements/title", headers=hx,
                        json={"lock_version": cn["lock_version"], "text": "침입"}).status_code == 403
    client.cookies.clear()
    assert client.get(f"/card-news/{cn['id']}/exports/{ok['id']}/download").status_code == 401
    from app.models.user import User

    owner = db_session.query(User).filter_by(email="p3@example.com").one()
    db_session.add(CardNewsExport(card_news_id=__import__("uuid").UUID(cn["id"]), file_format="pptx", status="running", created_by_user_id=owner.id))
    db_session.commit()
    h = auth_headers(client, "p3@example.com")
    assert client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "pptx"}).status_code == 409


def test_export_records_content_version_and_template_version(client, db_session):
    make_user(db_session, "p5@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "p5@example.com")
    cn = _fill_A(client, h, _create(client, h, "A", body_count=1))
    e1 = client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "pptx"}).json()
    assert e1["content_version"] == cn["lock_version"] and e1["template_version"] == 1 and e1["approval_state"] == "draft"
    changed = _set(client, h, cn, 1, "heading", text="수정 후 제목")
    e2 = client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "pptx"}).json()
    assert e2["content_version"] == changed["lock_version"] > e1["content_version"]
    listed = client.get(f"/card-news/{cn['id']}/exports", headers=h).json()
    assert len(listed) == 2  # 이전 출력본이 보존된다


def test_template_new_version_admin_only_and_applied(client, db_session):
    make_user(db_session, "a1@example.com", UserRole.ADMIN)
    make_user(db_session, "a2@example.com", UserRole.RESEARCHER)
    ha, hr = auth_headers(client, "a1@example.com"), auth_headers(client, "a2@example.com")
    assert client.post("/card-news/templates/A/new-version", headers=hr, json={"config": {"navy": "000000"}}).status_code == 403
    new = client.post("/card-news/templates/A/new-version", headers=ha, json={"config": {"navy": "102040", "bogus": "x"}})
    assert new.status_code == 201 and new.json()["version"] == 2 and new.json()["config"]["navy"] == "102040"
    templates = {t["key"]: t for t in client.get("/card-news/templates", headers=ha).json()}
    assert templates["A"]["version"] == 2 and "bogus" not in templates["A"]["config"]


# ---------------------------------------------------------------- PNG(실제 렌더링) ----
@requires_soffice
@pytest.mark.render
def test_png_export_renders_real_high_resolution_cards_same_ratio_as_pptx(client, db_session):
    make_user(db_session, "g1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "g1@example.com")
    cn = _fill_A(client, h, _create(client, h, "A", body_count=1))
    exp = client.post(f"/card-news/{cn['id']}/exports", headers=h, json={"file_format": "png", "short_px": 1000})
    assert exp.status_code == 201, exp.text
    e = exp.json()
    assert e["status"] == "success", e["error_message"]
    assert e["png_width_px"] >= 1000 and e["card_count"] == 3  # 짧은 변 1,000px 이상
    validation = json.loads(e["validation_json"])
    assert validation["rendered"] is True
    assert abs(validation["aspect_ratio_png"] - validation["aspect_ratio_pptx"]) < 0.01  # PPTX와 같은 비율
    zf = zipfile.ZipFile(io.BytesIO(client.get(f"/card-news/{cn['id']}/exports/{e['id']}/download", headers=h).content))
    names = sorted(zf.namelist())
    assert len(names) == 3 and all(n.endswith(".png") for n in names)
    from PIL import Image

    for name in names:
        img = Image.open(io.BytesIO(zf.read(name)))
        assert min(img.size) >= 1000 and abs(img.size[0] / img.size[1] - 7.0 / 15.62) < 0.01
    sample_dir = os.environ.get("REPORT_SAMPLE_DIR")  # CI가 합성 데이터 PNG를 아티팩트로 보관할 때 사용
    if sample_dir:
        from pathlib import Path

        out = Path(sample_dir) / "card_news"
        out.mkdir(parents=True, exist_ok=True)
        for name in names:
            (out / name).write_bytes(zf.read(name))

    one = client.get(f"/card-news/{cn['id']}/exports/{e['id']}/cards/2.png", headers=h)
    assert one.status_code == 200 and one.content[:8] == b"\x89PNG\r\n\x1a\n"
    bad = [i["message"] for i in validation["issues"] if i["code"] in ("rendered_out_of_bounds", "page_count_mismatch")]
    assert not bad, bad

