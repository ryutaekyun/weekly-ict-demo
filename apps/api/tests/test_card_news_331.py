"""Phase 3-3.1 — 외부 검토(카드뉴스 PPTX 4종·PNG 6장)에서 발견한 오류의 재발 방지 테스트.

- A형 빈 항목(번호 '3'과 빈 상자) 출력, 카드 번호 일관성, 앞머리 [..]·영문 약어 줄바꿈
- 출처 하이퍼링크·슬라이드 노트 출처 목록, 합성 테스트 데이터 표시와 확정 차단, 장식 도형과 실제 이탈 구분
- 후속 규정(시행일 연기 등) 알림 → 원고·보고서·카드뉴스 재검토 표시 → 연구원 확인
"""

import io
import json
import zipfile

import pytest
from pptx import Presentation

from app.models.enums import UserRole
from app.services import card_news_pptx, card_news_scene
from tests.conftest import auth_headers, make_user
from tests.test_card_news import (  # noqa: F401  (autouse 픽스처 재사용)
    _create, _el, _export_dir, _report_with_manuscript, _scene_for, _set, _skip_rendering_unless_marked, _t,
)


def _labels(scene):
    return [i["text"] for i in scene["items"] if i["type"] == "text"]


# ---------------------------------------------------------------- A형 빈 항목 ----
def test_template_a_hides_empty_content_boxes_and_renumbers_on_export():
    els = {"heading": _t("제목"), "content_1": _t("첫째 내용"), "content_2": _t("둘째 내용"), "content_3": _t(""), "source": _t("출처")}
    scene, cfg = _scene_for("A", "body", els)
    labels = _labels(scene)
    assert "1" in labels and "2" in labels and "3" not in labels  # 번호 3과 빈 상자가 출력에 없다
    boxes = [i for i in scene["items"] if i["type"] == "rect" and i.get("radius") == 0.18]
    assert len(boxes) == 2 and boxes[0]["h"] > 3.1  # 남은 상자는 높이를 키워 균형 있게 배치한다
    assert not [i for i in card_news_scene.check_scene(scene, cfg) if i["level"] == "error"]
    one, _ = _scene_for("A", "body", {**els, "content_2": _t("")})
    assert [i["text"] for i in one["items"] if i["type"] == "text" and i["text"] in ("1", "2", "3")] == ["1"]
    # 가운데 항목만 비면 남은 항목이 1, 2로 다시 매겨진다
    gap, _ = _scene_for("A", "body", {**els, "content_2": _t(""), "content_3": _t("셋째 내용")})
    assert [i["text"] for i in gap["items"] if i["type"] == "text" and i["text"] in ("1", "2", "3")] == ["1", "2"]


def test_template_a_editor_preview_still_shows_all_three_slots():
    els = {"heading": _t("제목"), "content_1": _t("첫째"), "content_2": _t(""), "content_3": _t(""), "source": _t("출처")}
    scene, _ = _scene_for("A", "body", els, for_export=False)
    assert {"1", "2", "3"} <= set(_labels(scene))  # 편집 화면은 입력 위치를 계속 보여준다


# ---------------------------------------------------------------- 번호·줄바꿈 ----
@pytest.mark.parametrize("template", ["A", "B", "C", "D"])
def test_cover_and_closing_have_same_card_number_tag_as_body(template):
    cover, _ = _scene_for(template, "cover", {"title": _t("표지 제목"), "source": _t("출처")}, )
    closing, _ = _scene_for(template, "closing", {"message": _t("메시지"), "source": _t("출처")})
    for scene in (cover, closing):
        assert "02/04" in _labels(scene)


def test_bracket_group_and_ascii_tokens_are_not_split_but_never_forced_when_too_wide():
    fits = card_news_scene.keep_bracket_together("[합성 테스트] 제목", 10, 5.9)
    assert fits.startswith("[합성 테스트]") and fits.endswith(" 제목")
    wide = card_news_scene.keep_bracket_together("[합성 테스트 데이터] 전기차 안전기준", 19, 5.9, True)
    assert " " not in wide  # 한 줄에 안 들어가는 묶음은 글자 중간에서 끊기지 않도록 그대로 둔다
    assert card_news_scene.keep_ascii_tokens("미국 FMVSS 305a 제정, No. 305a") == "미국 FMVSS 305a 제정, No. 305a"
    assert card_news_scene.keep_ascii_tokens("한글 단어 사이는 그대로") == "한글 단어 사이는 그대로"


# ---------------------------------------------------------------- 합성 표시·장식 ----
def test_synthetic_cards_carry_a_strong_banner_and_cannot_be_approved(client, db_session):
    scene, _ = _scene_for("C", "body", {"heading": _t("제목")}, is_synthetic=True, is_pilot=True)
    assert any("합성 테스트 데이터 · 실제 정책·통계 아님" in t for t in _labels(scene))
    plain, _ = _scene_for("C", "body", {"heading": _t("제목")})
    assert not any("합성 테스트 데이터" in t for t in _labels(plain))

    make_user(db_session, "s1@example.com", UserRole.RESEARCHER)
    make_user(db_session, "s2@example.com", UserRole.APPROVER)
    ha, hb = auth_headers(client, "s1@example.com"), auth_headers(client, "s2@example.com")
    cn = _create(client, ha, "C", "합성 카드")
    from app.models.card_news import CardNews
    row = db_session.get(CardNews, __import__("uuid").UUID(cn["id"]))
    row.is_synthetic = True
    row.status = "review_requested"
    db_session.commit()
    resp = client.post(f"/card-news/{cn['id']}/approve", headers=hb, json={"exception_reason": "내부 점검용 예외"})
    assert resp.status_code == 422 and "합성" in json.dumps(resp.json(), ensure_ascii=False)  # 예외 사유가 있어도 확정 불가


def test_decorative_shapes_may_bleed_but_real_shapes_leaving_the_card_are_errors():
    cover, cfg = _scene_for("C", "cover", {"title": _t("표지"), "source": _t("출처")})
    deco = [i for i in cover["items"] if i["type"] == "ellipse"]
    assert deco and all(i["deco"] for i in deco)
    assert any(i["x"] + i["w"] > cover["width"] for i in deco)  # 배경 장식은 카드 가장자리를 넘는다
    assert not [i for i in card_news_scene.check_scene(cover, cfg) if i["code"] == "out_of_bounds"]
    cover["items"].append({"type": "rect", "x": 5.0, "y": 1.0, "w": 4.0, "h": 1.0, "fill": "FFFFFF", "line": None, "radius": 0, "name": "box"})
    assert [i for i in card_news_scene.check_scene(cover, cfg) if i["code"] == "out_of_bounds"]


# ---------------------------------------------------------------- 출처 링크 ----
def _pptx_with_links(color):
    cfg = card_news_pptx  # noqa: F841
    scene, _ = _scene_for("A", "body", {"heading": _t("제목"), "content_1": _t("내용"), "source": _t("출처")},
                          )
    for it in scene["items"]:
        if it.get("key") == "source":
            it["links"] = [{"label": "NHTSA 2024-28707", "url": "https://www.federalregister.gov/d/2024-28707"},
                           {"label": "2025-02582", "url": "https://www.federalregister.gov/d/2025-02582"}]
            it["text"] = "출처: NHTSA 2024-28707, 2025-02582"
            it["color"] = color
    scene["sources"] = [{"label": "NHTSA 2024-28707", "url": "https://www.federalregister.gov/d/2024-28707", "doc_title": "최종 규정",
                         "published": "2024-12-20", "state": "verified"},
                        {"label": "NHTSA 2025-02582", "url": "https://www.federalregister.gov/d/2025-02582", "doc_title": "시행일 연기",
                         "published": "2025-02-14", "state": "not_verified"}]
    return card_news_pptx.build_pptx([scene], "테스트", "맑은 고딕")


def test_pptx_source_text_is_a_real_hyperlink_and_notes_list_every_source_with_status():
    prs = Presentation(io.BytesIO(_pptx_with_links("566070")))
    slide = prs.slides[0]
    source = next(s for s in slide.shapes if s.name == "source")
    runs = [r for p in source.text_frame.paragraphs for r in p.runs]
    urls = {r.hyperlink.address for r in runs if r.hyperlink.address}
    assert urls == {"https://www.federalregister.gov/d/2024-28707", "https://www.federalregister.gov/d/2025-02582"}
    notes = slide.notes_slide.notes_text_frame.text
    assert "2024-12-20" in notes and "연구원 검증됨" in notes and "근거 미검증" in notes and "2025-02582" in notes
    # PPTX는 여전히 편집 가능한 객체: 이미지 없음
    assert not [s for s in slide.shapes if s.shape_type == 13]


def test_pptx_on_dark_background_keeps_text_color_and_links_the_whole_box():
    prs = Presentation(io.BytesIO(_pptx_with_links("D5DDEC")))
    source = next(s for s in prs.slides[0].shapes if s.name == "source")
    assert source.click_action.hyperlink.address == "https://www.federalregister.gov/d/2024-28707"
    assert not any(r.hyperlink.address for p in source.text_frame.paragraphs for r in p.runs)  # 글자색이 테마 링크색으로 바뀌지 않는다


def test_source_links_are_shortened_to_one_line_with_remainder_count():
    links = [{"label": "NHTSA 2024-28707", "url": "u1"}, {"label": "NHTSA 2025-02582", "url": "u2"}, {"label": "EPA 2026-00001", "url": "u3"},
             {"label": "KATS 2026-00002", "url": "u4"}]
    shown, text = card_news_scene.fit_source_links(links, 7, 5.9)
    assert shown[0]["label"] == "NHTSA 2024-28707" and shown[1]["label"] == "2025-02582"  # 같은 기관 반복은 번호만
    assert text.startswith("출처: NHTSA 2024-28707, 2025-02582") and "외" in text
    assert card_news_scene.text_lines(text, 7, 5.9) == 1


# ---------------------------------------------------------------- 후속 규정 알림 ----
def _notice(client, h, **over):
    body = {"title": "시행일 연기", "original_doc_number": "2024-28707", "original_url": "https://www.federalregister.gov/documents/test",
            "followup_doc_number": "2025-02582", "followup_url": "https://www.federalregister.gov/d/2025-02582",
            "followup_org": "NHTSA", "followup_effective_on": "2025-03-20", "change_summary": "시행일을 2025.3.20.으로 연기",
            "changes": [{"item": "시행일", "old": "2025.2.18.", "new": "2025.3.20."}]}
    body.update(over)
    return client.post("/followups", headers=h, json=body)


def test_followup_flags_manuscript_report_and_card_news_without_changing_anything_then_ack_clears(client, db_session):
    make_user(db_session, "f1@example.com", UserRole.RESEARCHER)
    make_user(db_session, "f2@example.com", UserRole.RESEARCHER)
    h, other = auth_headers(client, "f1@example.com"), auth_headers(client, "f2@example.com")
    mid, rid = _report_with_manuscript(client, h, verify=True)
    sent = next(s for s in client.get(f"/card-news/report-sentences/{rid}", headers=h).json()["sentences"] if "제정안을 공고함" in s["sentence"])
    cn = client.post("/card-news", headers=h, json={"title": "후속 규정 연계", "template_key": "A", "source_type": "report",
                                                      "source_report_id": rid, "selections": [{"manuscript_id": mid, "sentence": sent["sentence"]}]}).json()
    assert client.get(f"/followups/for/card_news/{cn['id']}", headers=h).json()["open"] == []
    manuscript_before = client.get(f"/manuscripts/{mid}", headers=h).json()
    card_before = client.get(f"/card-news/{cn['id']}", headers=h).json()

    created = _notice(client, h)
    assert created.status_code == 201, created.text
    nid = created.json()["id"]
    impact = created.json()["impact"]
    assert [r["id"] for r in impact["manuscript"]] == [mid] and impact["report"][0]["id"] == rid and impact["card_news"][0]["id"] == cn["id"]

    # 표시만 하고 내용·검증 상태는 건드리지 않는다
    assert client.get(f"/manuscripts/{mid}", headers=h).json()["body"] == manuscript_before["body"]
    after = client.get(f"/card-news/{cn['id']}", headers=h).json()
    assert [e["text"] for c in after["cards"] for e in c["elements"]] == [e["text"] for c in card_before["cards"] for e in c["elements"]]
    assert after["followups_open"] and after["followups_open"][0]["followup_effective_on"] == "2025-03-20"
    blocking_codes = {b["code"] for b in after["blocking"]}
    assert "followup_pending" in blocking_codes  # 연구원 확인 전에는 내부 확정이 막힌다
    rep = client.get(f"/reports/{rid}", headers=h).json()
    assert any(b["code"] == "followup_pending" for b in rep["blocking"])

    # 권한·검증
    ack = {"entity_type": "card_news", "entity_id": cn["id"], "resolution": "updated", "note": "카드 내용을 정정하고 근거를 다시 확인함"}
    assert client.post(f"/followups/{nid}/ack", headers=other, json=ack).status_code == 403  # 남의 산출물은 확인 불가
    assert client.post(f"/followups/{nid}/ack", headers=h, json={**ack, "note": ""}).status_code == 422  # 사유 필수
    unrelated = _create(client, h, "A", "무관한 카드")
    assert client.post(f"/followups/{nid}/ack", headers=h, json={**ack, "entity_id": unrelated["id"]}).status_code == 422  # 영향 대상이 아님
    assert client.post(f"/followups/{nid}/ack", headers=h, json=ack).status_code == 200
    assert client.post(f"/followups/{nid}/ack", headers=h, json=ack).status_code == 409  # 중복 확인 불가
    assert client.get(f"/followups/for/card_news/{cn['id']}", headers=h).json()["open"] == []
    assert not any(b["code"] == "followup_pending" for b in client.get(f"/card-news/{cn['id']}", headers=h).json()["blocking"])
    # 다른 산출물은 아직 열려 있다(개별 확인)
    assert client.get(f"/followups/for/manuscript/{mid}", headers=h).json()["open"]
    assert client.get(f"/followups/for/report/{rid}", headers=h).json()["open"]
    listing = client.get("/followups", headers=h).json()
    assert listing[0]["open_count"] == 2 and listing[0]["affected_count"] == 3


def test_followup_requires_login_and_validates_input(client, db_session):
    assert client.get("/followups").status_code == 401
    make_user(db_session, "f3@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "f3@example.com")
    assert _notice(client, h, change_summary="").status_code == 422
    assert _notice(client, h, followup_url="x").status_code == 422
    made = _notice(client, h, original_doc_number="9999-00000").json()
    assert made["impact"] == {"manuscript": [], "report": [], "card_news": []}  # 영향받는 산출물이 없으면 아무것도 표시하지 않는다


def test_followup_candidates_endpoint_does_not_register_anything_automatically(client, db_session, monkeypatch):
    from app.services import followup_service
    make_user(db_session, "f4@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "f4@example.com")
    monkeypatch.setattr(followup_service, "detect_candidates", lambda db, timeout=20.0: [{"followup_doc_number": "2025-02582"}])
    assert client.post("/followups/detect", headers=h).json()["candidates"] == [{"followup_doc_number": "2025-02582"}]
    assert client.get("/followups", headers=h).json() == []  # 후보는 연구원이 확인해 등록하기 전까지 알림이 아니다
