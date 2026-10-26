"""Phase 3-4 — 주간기술동향지(외부 전문가 원고 의뢰·접수·검토·심의·교정·발간) 테스트.

모든 전문가·원고·계정은 합성 데이터이며 실제 인물·기관과 무관하다. 외부 이메일 발송·IITP 업로드는 구현하지 않으므로 테스트하지 않는다.
"""

import io
import json
import zipfile
from datetime import date

import pytest
from docx import Document

from app.core.config import settings
from app.models.enums import UserRole
from app.services import journal_service as jsvc
from app.services import report_pdf
from tests.conftest import auth_headers, make_user

_has_soffice = report_pdf.find_soffice() is not None
requires_soffice = pytest.mark.skipif(not _has_soffice, reason="LibreOffice 없음 — CI에서 검증")


@pytest.fixture(autouse=True)
def _dirs(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "journal_dir", str(tmp_path / "journal"))
    monkeypatch.setattr(settings, "export_dir", str(tmp_path / "exports"))


@pytest.fixture(autouse=True)
def _skip_rendering_unless_marked(request, monkeypatch):
    if request.node.get_closest_marker("render") is None:
        monkeypatch.setattr(report_pdf, "find_soffice", lambda: None)


def make_docx(title="전기차 안전 규제 동향", body=None, table=False, refs=True, macro=False) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_heading("1. 개요", level=2)
    for text in body or ["합성 테스트 문단 하나입니다.", "합성 테스트 문단 둘입니다. 실제 내용이 아닙니다."]:
        doc.add_paragraph(text)
    if table:
        doc.add_paragraph("[표 1] 합성 비교표")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "구분", "내용", "A", "합성 값"
    if refs:
        doc.add_heading("참고문헌", level=2)
        doc.add_paragraph("[1] 합성 참고문헌.")
    buf = io.BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    if macro:
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                zout.writestr(item, zin.read(item.filename))
            zout.writestr("word/vbaProject.bin", b"\x00fake-macro")
        data = out.getvalue()
    return data


def docx_file(name="원고.docx", **kw):
    return {"file": (name, make_docx(**kw), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}


class World:
    """관리자·편집자·승인권자·내부 검토자·외부 저자·외부 검토자를 한 번에 준비하는 테스트 도우미."""

    def __init__(self, client, db):
        self.c = client
        make_user(db, "adm@example.com", UserRole.ADMIN, "관리자")
        make_user(db, "ed@example.com", UserRole.RESEARCHER, "편집연구원")
        make_user(db, "ap@example.com", UserRole.APPROVER, "승인권자")
        make_user(db, "rv@example.com", UserRole.REVIEWER, "내부검토자")
        self.adm, self.ed, self.ap, self.rv = (auth_headers(client, e) for e in ("adm@example.com", "ed@example.com", "ap@example.com", "rv@example.com"))

    def expert_with_account(self, name="합성전문가", aff="가상대학교", email="author1@example.com"):
        e = self.c.post("/journal/experts", headers=self.ed, json={"name": name, "affiliation": aff, "email": email, "fields": ["AI"], "is_synthetic": True}).json()
        acc = self.c.post("/journal/accounts/author", headers=self.adm, json={"email": email, "display_name": name, "password": "TestPass!123", "expert_id": e["id"]})
        assert acc.status_code == 201, acc.text
        return e, auth_headers(self.c, email)

    def issue(self, no=1, publish="2027-03-31"):
        self.c.post("/journal/issues/plan-year", headers=self.ed, json={"year": 2027, "count": 3, "is_synthetic": True})
        issues = self.c.get("/journal/issues?year=2027", headers=self.ed).json()
        iss = next(i for i in issues if i["issue_no"] == no)
        if publish:
            r = self.c.patch(f"/journal/issues/{iss['id']}", headers=self.adm, json={"publish_date": publish})
            assert r.status_code == 200, r.text
        return iss["id"]

    def accepted_commission(self, issue_id, expert, topic="합성 원고 주제", due="2027-02-20"):
        c = self.c.post("/journal/commissions", headers=self.ed, json={"issue_id": issue_id, "expert_id": expert["id"], "topic": topic, "due_date": due}).json()
        for st in ("requested", "accepted"):
            r = self.c.post(f"/journal/commissions/{c['id']}/status", headers=self.ed, json={"status": st})
            assert r.status_code == 200, r.text
        return r.json()

    def submit(self, author_h, commission, **kw):
        title_form = kw.pop("title_form", "합성 원고 제목")
        files = docx_file(**kw)
        data = {"title": title_form, "title_en": "Synthetic Manuscript", "summary": "합성 요약",
                "authors_json": json.dumps([{"name": "합성전문가", "affiliation": "가상대학교", "corresponding": True}]),
                "references_text": "[1] 합성", "copyright_confirmed": "true"}
        return self.c.post(f"/journal/portal/author/commissions/{commission['id']}/submit", headers=author_h, files=files, data=data)


@pytest.fixture()
def w(client, db_session):
    return World(client, db_session)


# ---------------------------------------------------------------- 접근 제한 ----
def test_external_accounts_cannot_use_internal_apis(client, db_session, w):
    _, author_h = w.expert_with_account()
    assert client.get("/candidates", headers=author_h).status_code == 403
    assert client.get("/journal/issues", headers=author_h).status_code == 403
    assert client.get("/users", headers=author_h).status_code == 403
    assert client.get("/journal/portal/author/commissions", headers=author_h).status_code == 200  # 자신의 포털은 사용 가능
    assert client.get("/auth/me", headers=author_h).status_code == 200
    # 외부 계정은 일반 사용자 목록에 노출되지 않는다
    users = client.get("/users", headers=w.ed).json()
    assert all(u["role"] not in ("author", "ext_reviewer") for u in users)
    # 내부 검토자는 편집 API 불가(배정된 원고는 포털로만)
    assert client.get("/journal/manuscripts", headers=w.rv).status_code == 403
    assert client.post("/journal/accounts/author", headers=w.ed, json={"email": "x@example.com", "display_name": "x", "password": "TestPass!123", "expert_id": "00000000-0000-0000-0000-000000000000"}).status_code == 403


# ---------------------------------------------------------------- 회차·일정 ----
def test_year_plan_has_no_auto_dates_and_admin_date_drives_deadlines(client, w):
    r = client.post("/journal/issues/plan-year", headers=w.ed, json={"year": 2027, "count": 25})
    assert r.json() == {"created": 25, "total": 25}
    assert client.post("/journal/issues/plan-year", headers=w.ed, json={"year": 2027, "count": 25}).json()["created"] == 0  # 중복 생성 없음
    issues = client.get("/journal/issues?year=2027", headers=w.ed).json()
    assert len(issues) == 25 and all(i["publish_date"] is None for i in issues)  # 발간일을 임의로 확정하지 않는다
    one = client.get(f"/journal/issues/{issues[0]['id']}", headers=w.ed).json()
    assert all(s["due_date"] is None and s["state"] == "no_date" for s in one["steps"])
    assert client.patch(f"/journal/issues/{one['id']}", headers=w.ed, json={"publish_date": "2027-03-31"}).status_code == 403  # 관리자만 발간일 입력
    done = client.patch(f"/journal/issues/{one['id']}", headers=w.adm, json={"publish_date": "2027-03-31"}).json()
    by_key = {s["step_key"]: s for s in done["steps"]}
    assert by_key["publish"]["due_date"] == "2027-03-31" and by_key["iitp_check"]["due_date"] == "2027-03-28" or by_key["iitp_check"]["due_date"] == "2027-03-26"
    assert by_key["plan"]["window_start_date"] is not None and by_key["plan"]["window_start_days"] == 60
    # 발간일을 바꾸면 마감일이 다시 계산된다
    moved = client.patch(f"/journal/issues/{one['id']}", headers=w.adm, json={"publish_date": "2027-04-14"}).json()
    assert {s["step_key"]: s["due_date"] for s in moved["steps"]}["publish"] == "2027-04-14"
    assert {s["step_key"]: s["due_date"] for s in moved["steps"]}["receive"] != by_key["receive"]["due_date"]


def test_business_day_basis_and_holiday_adjustment():
    cfg = dict(jsvc.DEFAULT_SETTINGS)
    publish = date(2027, 3, 31)  # 수요일
    cal = jsvc.compute_due(publish, 3, {**cfg, "day_basis": "calendar", "holiday_rule": "none"})
    assert cal == date(2027, 3, 28)  # 일요일 그대로
    prev = jsvc.compute_due(publish, 3, {**cfg, "day_basis": "calendar", "holiday_rule": "prev"})
    assert prev == date(2027, 3, 26)  # 금요일로 앞당김
    nxt = jsvc.compute_due(publish, 3, {**cfg, "day_basis": "calendar", "holiday_rule": "next"})
    assert nxt == date(2027, 3, 29)
    biz = jsvc.compute_due(publish, 3, {**cfg, "day_basis": "business"})
    assert biz == date(2027, 3, 26)  # 영업일 3일 전: 3/30(화)·3/29(월)·3/26(금)
    holi = jsvc.compute_due(publish, 3, {**cfg, "day_basis": "business", "holidays": ["2027-03-29"]})
    assert holi == date(2027, 3, 25)  # 휴일(월) 제외


def test_overdue_steps_are_flagged_delayed_never_auto_completed(client, w, monkeypatch):
    iid = w.issue(publish="2027-03-31")
    monkeypatch.setattr(jsvc, "today", lambda: date(2027, 3, 10))
    issue = client.get(f"/journal/issues/{iid}", headers=w.ed).json()
    states = {s["step_key"]: s for s in issue["steps"]}
    assert states["receive"]["state"] == "delayed" and states["receive"]["completed_at"] is None  # 지연 표시, 자동 완료 아님
    assert "지연" in states["receive"]["impact"] and "다음 단계" in states["receive"]["impact"]
    sid = states["receive"]["id"]
    done = client.post(f"/journal/issues/{iid}/steps/{sid}/complete", headers=w.ed, json={"note": "실제 접수 완료"}).json()
    got = {s["step_key"]: s for s in done["steps"]}["receive"]
    assert got["state"] == "done" and got["due_date"] != got["completed_at"][:10] or got["state"] == "done"  # 마감일과 실제 완료일 구분
    dash = client.get("/journal/dashboard?year=2027", headers=w.ed).json()
    assert any(d["label"] for d in dash["delayed"])


def test_settings_are_admin_only_and_not_hardcoded(client, w):
    cfg = client.get("/journal/settings", headers=w.ed).json()
    assert cfg["word_min"] is None and cfg["word_max"] is None and cfg["required_docs"] == []  # 기준은 설정 전까지 비어 있다
    assert client.patch("/journal/settings", headers=w.ed, json={"word_max": 5000}).status_code == 403
    assert client.patch("/journal/settings", headers=w.adm, json={"word_max": 5000, "day_basis": "business"}).json()["word_max"] == 5000
    bad = client.patch("/journal/settings", headers=w.adm, json={"decision_outcomes": [{"key": "x", "label": "x", "effect": "weird"}]})
    assert bad.status_code == 422
    cleared = client.patch("/journal/settings", headers=w.adm, json={"clear": ["word_max"]}).json()
    assert cleared["word_max"] is None


# ---------------------------------------------------------------- 전문가·의뢰 ----
def test_expert_registry_duplicates_consent_privacy_and_anonymize(client, db_session, w):
    e = client.post("/journal/experts", headers=w.ed, json={"name": "김가상", "affiliation": "가상연구원", "email": "k@example.com", "fields": ["AI", "반도체"]}).json()
    assert e["expert_code"].startswith("E-") and e["fields"] == ["AI", "반도체"]
    dup = client.post("/journal/experts", headers=w.ed, json={"name": " 김가상 ", "affiliation": "가상 연구원", "email": "other@example.com"})
    assert dup.status_code == 409 and dup.json()["detail"]["existing_id"] == e["id"]  # 같은 사람 중복 등록 방지
    assert client.post("/journal/experts", headers=w.ed, json={"name": "다른이름", "affiliation": "다른기관", "email": "K@example.com"}).status_code == 409
    assert client.get("/journal/experts", headers=w.rv).status_code == 403  # 개인정보는 편집 담당자 이상만
    found = client.get("/journal/experts?field=AI", headers=w.ed).json()
    assert found[0]["name"] == "김가상" and "집필 가능 여부" in found[0]["note"]
    assert client.get("/journal/experts?q=없는이름", headers=w.ed).json() == []
    consent = client.post(f"/journal/experts/{e['id']}/consent", headers=w.ed, json={"consent_status": "agreed"}).json()
    assert consent["consent_status"] == "agreed" and consent["consent_at"]
    assert client.post(f"/journal/experts/{e['id']}/anonymize", headers=w.ed).status_code == 403
    anon = client.post(f"/journal/experts/{e['id']}/anonymize", headers=w.adm).json()
    assert anon["anonymized"] and anon["email"] is None and anon["affiliation"] == "(비식별)"
    iid = w.issue()
    assert client.post("/journal/commissions", headers=w.ed, json={"issue_id": iid, "expert_id": e["id"], "topic": "x"}).status_code == 409


def test_commission_flow_mail_is_never_sent_and_author_sees_only_own(client, w):
    e1, h1 = w.expert_with_account("저자하나", "가상대학교", "a1@example.com")
    e2, h2 = w.expert_with_account("저자둘", "가상연구소", "a2@example.com")
    iid = w.issue()
    c = client.post("/journal/commissions", headers=w.ed, json={"issue_id": iid, "expert_id": e1["id"], "topic": "합성 주제 1", "purpose": "취지",
                                                                  "due_date": "2027-02-20", "length_spec": "합성 분량"}).json()
    assert c["status"] == "draft" and c["mail_status"] == "none"
    prev = client.get(f"/journal/commissions/{c['id']}/mail-preview", headers=w.ed).json()
    assert "합성 주제 1" in prev["body"] and "발송되지 않" in prev["notice"]
    saved = client.put(f"/journal/commissions/{c['id']}/mail", headers=w.ed, json={"subject": prev["subject"], "body": prev["body"], "queue": True}).json()
    assert saved["mail_status"] == "pending_send" and saved["status"] == "ready" and "발송되지 않" in saved["mail_notice"]
    assert "sent" not in {saved["mail_status"], saved["status"]}  # '발송 완료' 상태는 존재하지 않는다
    # 아직 의뢰하지 않은 건은 저자에게 보이지 않는다
    assert client.get("/journal/portal/author/commissions", headers=h1).json() == []
    assert client.post(f"/journal/commissions/{c['id']}/status", headers=w.ed, json={"status": "accepted"}).status_code == 409  # 잘못된 전이
    client.post(f"/journal/commissions/{c['id']}/status", headers=w.ed, json={"status": "requested"})
    mine = client.get("/journal/portal/author/commissions", headers=h1).json()
    assert len(mine) == 1 and mine[0]["can_respond"] is True
    assert client.get("/journal/portal/author/commissions", headers=h2).json() == []  # 다른 저자는 볼 수 없다
    assert client.post(f"/journal/portal/author/commissions/{c['id']}/respond", headers=h2, json={"response": "accept"}).status_code == 404
    ok = client.post(f"/journal/portal/author/commissions/{c['id']}/respond", headers=h1, json={"response": "accept", "note": "수락합니다"})
    assert ok.status_code == 200 and ok.json()["status"] == "accepted"
    detail = client.get(f"/journal/commissions/{c['id']}", headers=w.ed).json()
    assert [e["to"] for e in detail["events"] if e["event"] == "status_change"] == ["ready", "requested", "accepted"]
    assert detail["manuscript_id"]  # 수락하면 원고 접수 건이 만들어진다
    # 미응답 → 늦은 수락
    c2 = client.post("/journal/commissions", headers=w.ed, json={"issue_id": iid, "expert_id": e2["id"], "topic": "합성 주제 2"}).json()
    for st in ("requested", "no_response", "accepted"):
        assert client.post(f"/journal/commissions/{c2['id']}/status", headers=w.ed, json={"status": st}).status_code == 200


# ---------------------------------------------------------------- 접수·보안 ----
def test_submission_versions_security_scan_and_format_check(client, w):
    e, author_h = w.expert_with_account()
    iid = w.issue()
    com = w.accepted_commission(iid, e)
    cid = com["id"]
    macro = w.submit(author_h, {"id": cid}, macro=True)
    assert macro.status_code == 422 and "매크로" in macro.json()["detail"]
    bad_ext = client.post(f"/journal/portal/author/commissions/{cid}/submit", headers=author_h, files={"file": ("a.exe", b"MZ....", "application/octet-stream")},
                          data={"title": "t", "authors_json": "[]"})
    assert bad_ext.status_code == 422 and "허용되지 않는" in bad_ext.json()["detail"]
    fake_pdf = client.post(f"/journal/portal/author/commissions/{cid}/submit", headers=author_h, files={"file": ("a.pdf", b"not a pdf", "application/pdf")},
                           data={"title": "t", "authors_json": "[]"})
    assert fake_pdf.status_code == 422
    js_pdf = client.post(f"/journal/portal/author/commissions/{cid}/submit", headers=author_h,
                         files={"file": ("a.pdf", b"%PDF-1.4\n1 0 obj<</Type/Catalog/OpenAction<</S/JavaScript/JS(app.alert(1))>>>>endobj", "application/pdf")},
                         data={"title": "t", "authors_json": "[]"})
    assert js_pdf.status_code == 422 and "스크립트" in js_pdf.json()["detail"]
    assert client.get("/journal/manuscripts", headers=w.ed).json()[0]["version_count"] == 0  # 거부된 파일은 접수되지 않는다

    ok = w.submit(author_h, {"id": cid}, table=True)
    assert ok.status_code == 201, ok.text
    ms = ok.json()
    assert ms["status"] == "received" and ms["version_count"] == 1 and ms["format_status"] == "ok"
    v = ms["versions"][0]
    assert len(v["sha256"]) == 64 and v["file_size"] > 0 and v["extract"]["tables"] == 1
    items = {i["key"]: i["status"] for i in v["format_check"]["items"]}
    assert items["authors"] == "pass" and items["references"] == "pass" and items["length"] == "info" and items["required_docs"] == "info"  # 기준 미설정
    assert client.post(f"/journal/portal/author/commissions/{cid}/submit", headers=author_h, files=docx_file(), data={"title": "t", "authors_json": "[]"}).status_code == 409
    # 접수 후 파일은 인증된 권한자만 내려받을 수 있다
    dl = client.get(f"/journal/manuscripts/{ms['id']}/versions/{v['id']}/download", headers=w.ed)
    assert dl.status_code == 200 and dl.content[:2] == b"PK" and "attachment" in dl.headers["content-disposition"]
    client.cookies.clear()  # 로그인 쿠키 없이는 내려받을 수 없다
    assert client.get(f"/journal/manuscripts/{ms['id']}/versions/{v['id']}/download").status_code == 401
    _, other_h = w.expert_with_account("다른저자", "다른대학", "other@example.com")
    assert client.get(f"/journal/portal/author/manuscripts/{ms['id']}", headers=other_h).status_code == 404
    assert client.get(f"/journal/portal/author/manuscripts/{ms['id']}/versions/{v['id']}/download", headers=other_h).status_code == 404
    assert client.get(f"/journal/portal/author/manuscripts/{ms['id']}/versions/{v['id']}/download", headers=author_h).status_code == 200


def test_required_docs_and_length_use_admin_configured_rules_only(client, w):
    client.patch("/journal/settings", headers=w.adm, json={"word_max": 5, "required_docs": [{"key": "consent", "label": "저작권 확인서"}]})
    e, author_h = w.expert_with_account()
    com = w.accepted_commission(w.issue(), e)
    ms = w.submit(author_h, com).json()
    assert ms["status"] == "format_fix_needed" and ms["format_status"] == "needs_fix"
    items = {i["key"]: i for i in ms["versions"][0]["format_check"]["items"]}
    assert items["length"]["status"] == "fail" and "저작권 확인서" in items["required_docs"]["detail"]
    # 제출서류 첨부 + 분량 기준 완화 후 보완본 재제출 → 별도 버전
    att = client.post(f"/journal/portal/author/manuscripts/{ms['id']}/attachments", headers=author_h, data={"doc_type": "consent"},
                      files={"file": ("확인서.pdf", b"%PDF-1.4\n%synthetic", "application/pdf")})
    assert att.status_code == 201, att.text
    client.patch("/journal/settings", headers=w.adm, json={"clear": ["word_max"]})
    again = w.submit(author_h, com)
    assert again.status_code == 201 and again.json()["version_count"] == 2 and again.json()["status"] == "received"
    assert len({v["sha256"] for v in again.json()["versions"]}) >= 1 and again.json()["versions"][0]["version_no"] == 1  # 이전 버전 보존


# ---------------------------------------------------------------- 검토·수정·심의 ----
def _to_in_review(w, author_h, expert):
    com = w.accepted_commission(w.issue(), expert)
    ms = w.submit(author_h, com).json()
    r = w.c.post(f"/journal/manuscripts/{ms['id']}/format-review", headers=w.ed, json={"result": "ok"})
    assert r.status_code == 200 and r.json()["status"] == "review_pending"
    return ms["id"], com


def test_review_revision_deliberation_full_flow_with_permissions(client, db_session, w):
    e, author_h = w.expert_with_account()
    mid, com = _to_in_review(w, author_h, e)
    rev = client.post("/journal/accounts/reviewer", headers=w.adm, json={"email": "extrev@example.com", "display_name": "외부검토자", "password": "TestPass!123"}).json()
    ext_h = auth_headers(client, "extrev@example.com")
    # 이해상충 경고: 같은 소속이면 확인 없이는 배정되지 않는다
    assert client.post(f"/journal/manuscripts/{mid}/assignments", headers=w.ed, json={"reviewer_user_id": rev["id"], "conflict_declared": True}).status_code == 409
    a = client.post(f"/journal/manuscripts/{mid}/assignments", headers=w.ed,
                    json={"reviewer_user_id": rev["id"], "due_date": "2027-03-01", "conflict_declared": False}).json()
    assert a["status"] == "assigned"
    assert client.get(f"/journal/manuscripts/{mid}", headers=w.ed).json()["status"] == "in_review"
    # 미배정 검토자(내부)는 접근 불가, 배정된 외부 검토자는 자신의 배정만 본다
    assert client.get(f"/journal/portal/reviews/{a['id']}", headers=w.rv).status_code == 404
    assert client.get("/journal/portal/reviews", headers=w.rv).json() == []
    mine = client.get("/journal/portal/reviews", headers=ext_h).json()
    assert len(mine) == 1 and mine[0]["id"] == a["id"]
    view = client.get(f"/journal/portal/reviews/{a['id']}", headers=ext_h).json()
    assert view["manuscript"]["title"] == "합성 원고 제목" and "참고 의견" in view["notice"]
    # 이해상충 확인 없이는 검토 완료 불가
    assert client.post(f"/journal/portal/reviews/{a['id']}/complete", headers=ext_h, json={"recommendation": "minor", "summary_opinion": "x"}).status_code == 409
    client.post(f"/journal/portal/reviews/{a['id']}/conflict", headers=ext_h, json={"declared": False})
    c1 = client.post(f"/journal/portal/reviews/{a['id']}/comments", headers=ext_h,
                     json={"scope": "paragraph", "anchor_text": "합성 테스트 문단 하나입니다.", "text": "근거를 보강해주세요."}).json()
    # 검토 의견은 문단·전체에 각각 연결할 수 있다
    client.post(f"/journal/portal/reviews/{a['id']}/comments", headers=ext_h, json={"scope": "doc", "text": "결론이 약합니다."})
    done = client.post(f"/journal/portal/reviews/{a['id']}/complete", headers=ext_h, json={"recommendation": "minor", "summary_opinion": "경미한 수정 필요"})
    assert done.status_code == 200 and done.json()["assignment"]["status"] == "completed"
    # 수정 요청 → 저자는 검토자 이름 없이 의견을 본다
    rr = client.post(f"/journal/manuscripts/{mid}/request-revision", headers=w.ed, json={})
    assert rr.status_code == 200 and rr.json()["status"] == "revision_requested"
    author_view = client.get(f"/journal/portal/author/manuscripts/{mid}", headers=author_h).json()
    assert {c["author_name"] for c in author_view["comments"]} == {"검토자"} and len(author_view["comments"]) == 2
    # 저자 답변(해결 아님) → 수정본 제출
    rep = client.post(f"/journal/portal/author/manuscripts/{mid}/comments/{c1['id']}/reply", headers=author_h, json={"text": "문단을 보강했습니다"}).json()
    assert rep["status"] == "answered" and rep["author_reply"]
    assert client.post(f"/journal/portal/author/manuscripts/{mid}/revision", headers=author_h, files=docx_file(body=["완전히 새로 쓴 문단입니다.", "보강한 두 번째 문단입니다."])).status_code == 201
    ms = client.get(f"/journal/manuscripts/{mid}", headers=w.ed).json()
    assert ms["status"] == "revision_received" and ms["version_count"] == 2 and ms["content_status"] == "pending"
    assert ms["review_validity"]["needs_rereview"] is True  # 변경 폭이 커서 이전 검토 유효성 재검토 필요
    # 버전 비교와 의견 대응 관계 — 수정 확정으로 단정하지 않는다
    cmp = client.get(f"/journal/manuscripts/{mid}/compare?from=1&to=2", headers=w.ed).json()
    assert cmp["counts"]["modified"] + cmp["counts"]["added"] + cmp["counts"]["deleted"] >= 1 and cmp["identical_text"] is False
    mapped = {m["id"]: m for m in cmp["comment_mapping"]}
    assert mapped[c1["id"]]["match"] == "changed" and "내용 확인 필요" in mapped[c1["id"]]["note"]
    assert "해결 처리는 담당자가 확인" in cmp["notice"]
    # 2차 검토 후 의견 해결·심의 결과 확정(승인권자만)
    a2 = client.post(f"/journal/manuscripts/{mid}/assignments", headers=w.ed, json={"reviewer_user_id": rev["id"], "round": 2}).json()
    assert client.get(f"/journal/manuscripts/{mid}", headers=w.ed).json()["status"] == "in_review"
    client.post(f"/journal/portal/reviews/{a2['id']}/conflict", headers=ext_h, json={"declared": False})
    for cm in client.get(f"/journal/manuscripts/{mid}", headers=w.ed).json()["comments"]:
        assert client.post(f"/journal/manuscripts/{mid}/comments/{cm['id']}/resolve", headers=w.ed).status_code == 200
    assert client.post(f"/journal/manuscripts/{mid}/deliberate", headers=w.ed).status_code == 409  # 2차 검토 미완료
    client.post(f"/journal/portal/reviews/{a2['id']}/complete", headers=ext_h, json={"recommendation": "accept", "summary_opinion": "수정 확인"})
    assert client.post(f"/journal/manuscripts/{mid}/deliberate", headers=w.ed).json()["status"] == "deliberated"
    assert client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ed, json={"outcome": "accept"}).status_code == 403  # 편집자는 게재 확정 불가
    assert client.post(f"/journal/manuscripts/{mid}/decision", headers=w.rv, json={"outcome": "accept"}).status_code == 403
    assert client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "nope"}).status_code == 422
    final = client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "accept", "note": "게재 확정"}).json()
    assert final["status"] == "accepted" and final["decisions"][-1]["is_final"] is True and final["decisions"][-1]["decided_by"] == "승인권자"


def test_accept_is_blocked_while_revision_comments_unresolved_and_minor_outcome_returns_for_revision(client, w):
    e, author_h = w.expert_with_account()
    mid, _ = _to_in_review(w, author_h, e)
    rev = client.post("/journal/accounts/reviewer", headers=w.adm, json={"email": "r2@example.com", "display_name": "검토자2", "password": "TestPass!123"}).json()
    ext_h = auth_headers(client, "r2@example.com")
    a = client.post(f"/journal/manuscripts/{mid}/assignments", headers=w.ed, json={"reviewer_user_id": rev["id"]}).json()
    client.post(f"/journal/portal/reviews/{a['id']}/conflict", headers=ext_h, json={"declared": False})
    client.post(f"/journal/portal/reviews/{a['id']}/comments", headers=ext_h, json={"text": "미해결 의견"})
    client.post(f"/journal/portal/reviews/{a['id']}/complete", headers=ext_h, json={"recommendation": "accept", "summary_opinion": "검토자는 적합 의견"})
    client.post(f"/journal/manuscripts/{mid}/deliberate", headers=w.ed)
    blocked = client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "accept"})
    assert blocked.status_code == 409 and "해결되지 않은" in blocked.json()["detail"]  # 검토자 추천만으로 게재 확정되지 않는다
    minor = client.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "minor", "note": "경미 수정"}).json()
    assert minor["status"] == "revision_requested" and minor["decisions"][-1]["is_final"] is False


def test_manuscript_lock_conflict_and_proof_to_issue_export(client, w):
    e, author_h = w.expert_with_account()
    mid, _ = _to_in_review(w, author_h, e)
    cur = client.get(f"/journal/manuscripts/{mid}", headers=w.ed).json()
    assert client.patch(f"/journal/manuscripts/{mid}", headers=w.ed, json={"lock_version": cur["lock_version"], "title_en": "A"}).status_code == 200
    assert client.patch(f"/journal/manuscripts/{mid}", headers=w.ed, json={"lock_version": cur["lock_version"], "title_en": "B"}).status_code == 409  # 동시 수정 충돌
    # 검토 없이 심의로 넘길 수 없다
    assert client.post(f"/journal/manuscripts/{mid}/deliberate", headers=w.ed).status_code == 409


def _accept(w, author_h, expert, issue_id=None, title_form="합성 원고 제목", body=None):
    c = w.c
    iid = issue_id or w.issue()
    com = w.accepted_commission(iid, expert, topic=title_form)
    ms = w.submit(author_h, com, title_form=title_form, **({"body": body} if body else {})).json()
    mid = ms["id"]
    c.post(f"/journal/manuscripts/{mid}/format-review", headers=w.ed, json={"result": "ok"})
    rev = c.get("/journal/reviewer-candidates", headers=w.ed).json()
    reviewer = next(r for r in rev if r["display_name"] == "내부검토자")
    a = c.post(f"/journal/manuscripts/{mid}/assignments", headers=w.ed, json={"reviewer_user_id": reviewer["id"]}).json()
    c.post(f"/journal/portal/reviews/{a['id']}/conflict", headers=w.rv, json={"declared": False})
    assert c.post(f"/journal/portal/reviews/{a['id']}/complete", headers=w.rv, json={"recommendation": "accept", "summary_opinion": "적합"}).status_code == 200
    c.post(f"/journal/manuscripts/{mid}/deliberate", headers=w.ed)
    assert c.post(f"/journal/manuscripts/{mid}/decision", headers=w.ap, json={"outcome": "accept"}).status_code == 200
    return mid, iid


def test_issue_composition_proof_gate_and_draft_vs_final_export(client, w):
    e1, h1 = w.expert_with_account("저자하나", "가상대학교", "p1@example.com")
    e2, h2 = w.expert_with_account("저자둘", "가상연구소", "p2@example.com")
    m1, iid = _accept(w, h1, e1, title_form="첫 번째 합성 원고", body=["첫 원고 본문입니다."])
    m2, _ = _accept(w, h2, e2, issue_id=iid, title_form="두 번째 합성 원고", body=["둘째 원고 본문입니다."])
    # 교정 확인 전에는 확정 발간본 불가, 초안 출력은 가능
    assert client.put(f"/journal/issues/{iid}/items", headers=w.ed, json={"manuscript_ids": [m2, m1]}).status_code == 200
    issue = client.get(f"/journal/issues/{iid}", headers=w.ed).json()
    assert [i["title"] for i in issue["items"]] == ["두 번째 합성 원고", "첫 번째 합성 원고"]  # 게재 순서 지정
    codes = {p["code"] for p in issue["export_preflight"]}
    assert "proof_not_confirmed" in codes and "iitp_unconfirmed" in codes
    assert client.post(f"/journal/issues/{iid}/exports", headers=w.ap, json={"file_format": "docx", "final": True}).status_code == 422
    assert client.post(f"/journal/issues/{iid}/exports", headers=w.ed, json={"file_format": "docx", "final": True}).status_code == 403
    draft = client.post(f"/journal/issues/{iid}/exports", headers=w.ed, json={"file_format": "docx"})
    assert draft.status_code == 201 and draft.json()["status"] == "success" and draft.json()["is_final"] is False
    # 교정본 업로드 → 점검표 전부 확인해야 최종 교정 확인
    for mid in (m1, m2):
        up = client.post(f"/journal/manuscripts/{mid}/versions", headers=w.ed, files=docx_file(), data={"kind": "proofread"})
        assert up.status_code == 201, up.text
        assert client.post(f"/journal/manuscripts/{mid}/proof/confirm", headers=w.ed, json={"checklist": {"source": True}}).status_code == 422
        full = {k: True for k in ("source", "figure_table_numbers", "references", "abbreviations", "language")}
        assert client.post(f"/journal/manuscripts/{mid}/proof/confirm", headers=w.ed, json={"checklist": full}).json()["proof_confirmed"] is True
    client.put(f"/journal/issues/{iid}/items", headers=w.ed, json={"manuscript_ids": [m2, m1]})
    final = client.post(f"/journal/issues/{iid}/exports", headers=w.ap, json={"file_format": "docx", "final": True})
    assert final.status_code == 201 and final.json()["status"] == "success" and final.json()["is_final"] is True
    fx = final.json()
    assert [i["version_no"] for i in fx["items"]] == [2, 2]  # 교정본 버전이 고정된다
    dl = client.get(f"/journal/issues/{iid}/exports/{fx['id']}/download", headers=w.ed)
    doc = Document(io.BytesIO(dl.content))
    texts = [p.text for p in doc.paragraphs]
    assert any("두 번째 합성 원고" in t for t in texts) and any("합성전문가(가상대학교)" in t for t in texts)  # 제목·저자·소속 표시
    assert texts.index(next(t for t in texts if t.startswith("1. 두 번째"))) < texts.index(next(t for t in texts if t.startswith("2. 첫 번째")))
    assert fx["validation"]["issues"] and any(i["code"] == "iitp_unconfirmed" for i in fx["validation"]["issues"])
    client.cookies.clear()
    assert client.get(f"/journal/issues/{iid}/exports/{fx['id']}/download").status_code == 401
    # 발간 기록과 IITP 확인은 별개 — 확인 전에는 '확인 완료'로 표시되지 않는다
    rec = client.post(f"/journal/issues/{iid}/publish", headers=w.ap).json()
    assert rec["status"] == "published" and rec["iitp_confirmed"] is False and rec["published_at"]
    assert client.patch(f"/journal/issues/{iid}", headers=w.ed, json={"status": "editing"}).status_code == 409
    assert client.put(f"/journal/issues/{iid}/items", headers=w.ed, json={"manuscript_ids": [m1]}).status_code == 409
    assert client.post(f"/journal/issues/{iid}/iitp", headers=w.ed, json={"confirmed": True}).status_code == 403
    assert client.post(f"/journal/issues/{iid}/iitp", headers=w.ap, json={"confirmed": True, "note": "실제 확인함"}).json()["iitp_confirmed"] is True


def test_unaccepted_or_pdf_manuscripts_cannot_be_composed_or_finalized(client, w):
    e, h = w.expert_with_account()
    com = w.accepted_commission(w.issue(), e)
    ms = w.submit(h, com).json()
    iid = ms["issue_id"]
    assert client.put(f"/journal/issues/{iid}/items", headers=w.ed, json={"manuscript_ids": [ms["id"]]}).status_code == 422  # 게재 확정 전


def test_dashboard_counts_come_from_database_not_hardcoded(client, w):
    client.post("/journal/issues/plan-year", headers=w.ed, json={"year": 2027, "count": 4})
    dash = client.get("/journal/dashboard?year=2027", headers=w.ed).json()
    assert dash["plan"]["planned_issues"] == 4 and dash["plan"]["published_records"] == 0  # 목표치 25를 흉내 내지 않는다
    assert dash["plan"]["with_publish_date"] == 0
    e, h = w.expert_with_account()
    iid = client.get("/journal/issues?year=2027", headers=w.ed).json()[0]["id"]
    w.accepted_commission(iid, e)
    d2 = client.get("/journal/dashboard?year=2027", headers=w.ed).json()
    assert d2["commission_status"] == {"accepted": 1}
    mine = client.get("/journal/dashboard?year=2027&mine=true", headers=w.ed).json()
    assert mine["issues"] == []  # 담당 회차가 없으면 비어 있다


def test_overdue_manuscript_and_review_deadlines_surface_as_delayed(client, w, monkeypatch):
    e, h = w.expert_with_account()
    iid = w.issue()
    w.accepted_commission(iid, e, due="2027-01-10")
    monkeypatch.setattr(jsvc, "today", lambda: date(2027, 1, 20))
    dash = client.get("/journal/dashboard?year=2027", headers=w.ed).json()
    assert any(d["kind"] == "manuscript_due" and d["state"] == "delayed" for d in dash["delayed"])


def test_template_versions_and_provisional_notice(client, w):
    t = client.get("/journal/template", headers=w.ed).json()
    assert t["is_provisional"] is True and "공식 주간기술동향지 양식" in t["notice"] and t["version"] == 1
    assert client.patch("/journal/template", headers=w.ed, json={"config": {"font_body": "Noto Sans CJK KR"}}).status_code == 403
    assert client.patch("/journal/template", headers=w.adm, json={"config": {"font_body": "Noto Sans CJK KR"}}).json()["version"] == 2


@requires_soffice
@pytest.mark.render
def test_pdf_export_renders_with_article_pages(client, w):
    e1, h1 = w.expert_with_account("저자하나", "가상대학교", "pp1@example.com")
    m1, iid = _accept(w, h1, e1, title_form="렌더링 합성 원고", body=["렌더링 확인용 본문입니다."])
    client.post(f"/journal/manuscripts/{m1}/versions", headers=w.ed, files=docx_file(), data={"kind": "proofread"})
    client.post(f"/journal/manuscripts/{m1}/proof/confirm", headers=w.ed, json={"checklist": {k: True for k in ("source", "figure_table_numbers", "references", "abbreviations", "language")}})
    client.put(f"/journal/issues/{iid}/items", headers=w.ed, json={"manuscript_ids": [m1]})
    pdf = client.post(f"/journal/issues/{iid}/exports", headers=w.ed, json={"file_format": "pdf"}).json()
    assert pdf["status"] == "success", pdf["error_message"]
    assert pdf["page_count"] >= 2 and pdf["validation"]["toc_verified"] is True
    assert pdf["validation"]["part_pages"][0]["start_page"] >= 1  # 원고별 시작 페이지 확인
