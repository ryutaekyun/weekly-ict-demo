"""Phase 3-2.1 — 외부 검토(시범 출력물 6종)에서 발견한 출력 품질 오류의 재발 방지 테스트.

- 동일한 소제목이 여러 원고에 반복될 때 목차 쪽수가 첫 번째 원고 것으로 잘못 나오던 오류
- 목차 항목 내부 이동 링크, 각주 URL의 실제 하이퍼링크, 재인용 축약
- 표 제목 입력, 확정본에 '미완성' 문구가 남던 모순, 간결형 요약본, 표지 제목의 영문 약어 줄바꿈
"""

import io
import re
import zipfile

import pytest
from docx import Document

from app.core.config import settings
from app.models.enums import UserRole
from app.services import report_docx, report_pdf
from tests.conftest import auth_headers, make_user
from tests.test_reports import (
    _body_doc, _cell, _general_report, _make_manuscript, _plain, _t, requires_soffice,
)


@pytest.fixture(autouse=True)
def _export_dir(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "export_dir", str(tmp_path / "exports"))


@pytest.fixture(autouse=True)
def _skip_rendering_unless_marked(request, monkeypatch):
    if request.node.get_closest_marker("render") is None:
        monkeypatch.setattr(report_pdf, "find_soffice", lambda: None)


def _snapshot(parts, **report):
    base = {"title": "테스트 보고서", "title_en": None, "kind_label": "일반호", "iitp_confirmed": False, "is_pilot": False,
            "status": "draft", "layout": "standard", "exception": False}
    base.update(report)
    return {"report": base, "template": {"name": "임시 표준 템플릿", "version": 1, "is_provisional": True, "config": {}}, "parts": parts}


def _fn(url="https://www.federalregister.gov/x", status="matched", location="3문단"):
    return {"source_org": "NHTSA", "doc_title": "FMVSS 305a", "source_url": url, "published_at": "2024-04-15T00:00:00+00:00",
            "accessed_at": "2026-10-23T00:00:00+00:00", "verification_status": status, "location_label": location}


def _doc_xml(docx_bytes):
    return zipfile.ZipFile(io.BytesIO(docx_bytes)).read("word/document.xml").decode("utf-8")


# ---------------------------------------------------------------- 목차 ----
def _toc_numbers(docx_bytes):
    xml = _doc_xml(docx_bytes)
    out = []
    for m in re.finditer(r'<w:hyperlink w:anchor="(_Toc\d+)" w:history="1">(.*?)</w:hyperlink>', xml, re.S):
        texts = re.findall(r"<w:t(?: [^>]*)?>(.*?)</w:t>", m.group(2))
        out.append((m.group(1), texts[0], int(texts[-1]) if len(texts) > 1 and texts[-1].isdigit() else None))
    return out


def test_toc_entries_are_linked_to_unique_bookmarks_even_when_titles_repeat():
    body = _body_doc("원고", "m")
    parts = [{"key": "key_issue", "heading": "Ⅰ. 주요 이슈", "manuscripts": [
        {"title": "원고 가", "body_json": body, "footnotes": {"m": _fn()}},
        {"title": "원고 나", "body_json": body, "footnotes": {"m": _fn()}}]}]
    built = report_docx.build_docx(_snapshot(parts))
    xml = _doc_xml(built.docx_bytes)
    sections = [h for h in built.headings if h[1] == "1. 추진 경과"]
    assert len(sections) == 2  # 같은 제목이 두 번
    anchors = [a for a, _t_, _n in _toc_numbers(built.docx_bytes)]
    assert len(anchors) == len(built.headings) == len(set(anchors))  # 항목마다 고유 북마크로 연결
    for name in anchors:
        assert f'w:name="{name}"' in xml  # 이동할 북마크가 실제로 있다


def test_toc_uses_page_per_entry_not_per_title_text():
    body = _body_doc("원고", "m")
    parts = [{"key": "k", "heading": "Ⅰ. 파트", "manuscripts": [
        {"title": "원고 가", "body_json": body, "footnotes": {"m": _fn()}},
        {"title": "원고 나", "body_json": body, "footnotes": {"m": _fn()}}]}]
    snap = _snapshot(parts)
    first = report_docx.build_docx(snap)
    pages = [3, 3, 3, 6, 6]  # 제목 5개(파트·원고·소제목 …) 순서 그대로의 쪽수
    assert len(first.headings) == len(pages)
    numbered = report_docx.build_docx(snap, pages)
    assert [n for _a, _t_, n in _toc_numbers(numbered.docx_bytes)] == pages  # 같은 제목이어도 항목별 번호가 들어간다


@requires_soffice
@pytest.mark.render
def test_rendered_pdf_toc_matches_real_pages_for_duplicate_subheadings_and_has_links(client, db_session):
    make_user(db_session, "t321a@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "t321a@example.com")
    m1 = _make_manuscript(client, h, title="원고 가")
    m2 = _make_manuscript(client, h, title="원고 나")
    rid = _general_report(client, h, {"key_issue": m1, "country_trend": m2})
    pdf = client.post(f"/reports/{rid}/exports", json={"file_format": "pdf"}, headers=h).json()
    assert pdf["status"] == "success", pdf["error_message"]
    assert '"toc_verified": true' in pdf["validation_json"]
    pdf_bytes = client.get(f"/reports/{rid}/exports/{pdf['id']}/download", headers=h).content
    docx = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h).json()
    docx_bytes = client.get(f"/reports/{rid}/exports/{docx['id']}/download", headers=h).content

    headings = [(1, "Ⅰ. 주요 이슈"), (2, "원고 가"), (3, "1. 추진 경과"), (1, "Ⅱ. 주요국 동향"), (2, "원고 나"), (3, "1. 추진 경과"),
                (1, "Ⅲ. ICT 9대 분야 주요 단신"), (1, "Ⅳ. 주요 일정")]
    actual = report_pdf.analyze_pdf(pdf_bytes, headings, [])["heading_page_list"]
    assert None not in actual and actual[2] != actual[5]  # 같은 소제목이 서로 다른 쪽에 있다
    toc = [n for _a, _txt, n in _toc_numbers(docx_bytes)]
    assert toc == actual  # 목차 번호가 실제 렌더링 쪽수와 항목마다 일치
    toc_page = report_pdf.analyze_pdf(pdf_bytes, headings, [])["toc_end_page"]
    assert report_pdf.count_internal_links(pdf_bytes, list(range(2, toc_page + 1))) >= len(headings)  # 목차→본문 내부 이동 링크


@requires_soffice
@pytest.mark.render
def test_docx_text_is_fully_present_in_rendered_pdf(client, db_session):
    make_user(db_session, "t321b@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "t321b@example.com")
    mid = _make_manuscript(client, h)
    rid = _general_report(client, h, {"key_issue": mid})
    docx = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h).json()
    pdf = client.post(f"/reports/{rid}/exports", json={"file_format": "pdf"}, headers=h).json()
    docx_bytes = client.get(f"/reports/{rid}/exports/{docx['id']}/download", headers=h).content
    pdf_text = "".join(report_pdf.analyze_pdf(client.get(f"/reports/{rid}/exports/{pdf['id']}/download", headers=h).content, [], [])["text"].split())
    missing = []
    for p in Document(io.BytesIO(docx_bytes)).paragraphs:
        probe = "".join(p.text.split())[:14]
        if probe and "\t" not in p.text and probe not in pdf_text:  # 목차 줄은 점선 지시선이 끼므로 제외
            missing.append(p.text)
    assert not missing, missing  # DOCX 본문이 PDF에 빠짐없이 들어 있다


# ---------------------------------------------------------------- 각주 ----
def test_footnote_url_is_real_word_hyperlink_and_repeat_citation_is_short():
    body = {"type": "doc", "content": [
        {"type": "paragraph", "content": [_t("첫 번째 근거 문장", "a"), _t(". 두 번째 근거 문장", "b")]}]}
    parts = [{"key": "k", "heading": "Ⅰ. 파트", "manuscripts": [
        {"title": "원고", "body_json": body, "footnotes": {"a": _fn(status="not_verified", location=None), "b": _fn()}}]}]
    built = report_docx.build_docx(_snapshot(parts))
    zf = zipfile.ZipFile(io.BytesIO(built.docx_bytes))
    notes = zf.read("word/footnotes.xml").decode("utf-8")
    rels = zf.read("word/_rels/footnotes.xml.rels").decode("utf-8")
    assert notes.count("<w:hyperlink") == 1  # 첫 인용에만 전체 서지정보와 링크
    assert "https://www.federalregister.gov/x" in rels and "TargetMode=\"External\"" in rels
    assert "앞의 각주 1번 참조" in notes  # 같은 출처 재인용은 축약
    assert notes.count("https://www.federalregister.gov/x") == 1
    assert "[근거 미검증]" in notes and "[원문 내 위치 미확인]" in notes  # 상태 표시는 각주마다 유지
    assert notes.count("<w:footnote ") - 2 == built.footnote_count == 2


# ---------------------------------------------------------------- 표 제목 ----
def test_table_caption_text_is_printed_after_the_number_and_first_row_kept_with_caption():
    table = {"type": "table", "attrs": {"caption": "국가별 주요 정책 비교"}, "content": [
        {"type": "tableRow", "content": [_cell("tableHeader", "국가"), _cell("tableHeader", "정책")]},
        {"type": "tableRow", "content": [_cell("tableCell", "A국"), _cell("tableCell", "내용")]}]}
    nocap = {"type": "table", "content": table["content"]}
    body = {"type": "doc", "content": [{"type": "paragraph", "content": [_t("본문")]}, table, nocap]}
    parts = [{"key": "k", "heading": "Ⅰ. 파트", "manuscripts": [{"title": "원고", "body_json": body, "footnotes": {}}]}]
    built = report_docx.build_docx(_snapshot(parts))
    texts = [p.text for p in Document(io.BytesIO(built.docx_bytes)).paragraphs]
    assert "[표 1-1] 국가별 주요 정책 비교" in texts and "[표 1-2]" in texts  # 제목이 없으면 번호만
    assert 'w:tblHeader' in _doc_xml(built.docx_bytes)


# ---------------------------------------------------------------- 표지 ----
def test_cover_title_keeps_acronyms_together_and_shrinks_long_titles():
    snap = _snapshot([], title="미국 NHTSA 전기차 동력계 안전기준 FMVSS 305a 제정 동향과 시행 일정 및 준수 시한 분석")
    built = report_docx.build_docx(snap)
    doc = Document(io.BytesIO(built.docx_bytes))
    cover = next(p for p in doc.paragraphs if p.style.name == "Title")
    assert "FMVSS 305a" in cover.text  # 영문 약어·번호가 줄바꿈에서 쪼개지지 않는다
    assert cover.runs[0].font.size.pt <= 19  # 긴 제목은 글자 크기를 줄인다
    texts = "\n".join(p.text for p in doc.paragraphs)
    assert "작성: 테크노베이션파트너스(TNP)" in texts and "문서 상태: 내부 검토용 초안" in texts


# ---------------------------------------------------------------- 확정본 정합성 ----
def _approve(client, ha, hb, mid):
    reviewer = next(u for u in client.get("/users", headers=ha).json() if u["role"] == "reviewer")
    assert client.post(f"/manuscripts/{mid}/review-request", headers=ha, json={"reviewer_user_id": reviewer["id"]}).status_code == 200
    assert client.post(f"/manuscripts/{mid}/approve", headers=hb, json={}).status_code == 200


def _docx_of(client, h, rid):
    exp = client.post(f"/reports/{rid}/exports", json={"file_format": "docx"}, headers=h).json()
    assert exp["status"] == "success", exp["error_message"]
    data = client.get(f"/reports/{rid}/exports/{exp['id']}/download", headers=h).content
    zf = zipfile.ZipFile(io.BytesIO(data))
    footers = "".join(zf.read(n).decode("utf-8") for n in zf.namelist() if n.startswith("word/footer"))
    return Document(io.BytesIO(data)), footers


def test_final_export_omits_empty_optional_parts_and_labels_status_consistently(client, db_session):
    make_user(db_session, "f1@example.com", UserRole.RESEARCHER)
    make_user(db_session, "f2@example.com", UserRole.APPROVER)
    make_user(db_session, "f3@example.com", UserRole.REVIEWER)
    ha, hb = auth_headers(client, "f1@example.com"), auth_headers(client, "f2@example.com")
    m1 = _make_manuscript(client, ha, title="주요 이슈 원고", verify=True)
    m2 = _make_manuscript(client, ha, title="주요국 원고", verify=True)
    for mid in (m1, m2):
        _approve(client, ha, hb, mid)
    rid = _general_report(client, ha, {"key_issue": m1, "country_trend": m2})

    draft_doc, _ = _docx_of(client, ha, rid)
    assert any("원고 미배치" in p.text for p in draft_doc.paragraphs)  # 초안은 비어 있는 선택 파트를 명시한다
    assert client.post(f"/reports/{rid}/finalize", json={}, headers=hb).status_code == 200
    final_doc, footers = _docx_of(client, ha, rid)
    texts = [p.text for p in final_doc.paragraphs]
    assert not any("원고 미배치" in t for t in texts)  # 확정본에 '미완성' 문구가 남지 않는다
    assert "Ⅲ. ICT 9대 분야 주요 단신" not in texts and "Ⅳ. 주요 일정" not in texts  # 빈 선택 파트는 출력 생략
    assert "Ⅰ. 주요 이슈" in texts and "Ⅱ. 주요국 동향" in texts
    assert "출력용 확정본(내부 절차)" in footers and "예외 승인" not in footers


def test_exception_approval_is_recorded_on_the_output_when_required_part_is_missing(client, db_session):
    make_user(db_session, "g1@example.com", UserRole.RESEARCHER)
    make_user(db_session, "g2@example.com", UserRole.APPROVER)
    ha, hb = auth_headers(client, "g1@example.com"), auth_headers(client, "g2@example.com")
    mid = _make_manuscript(client, ha)
    rid = _general_report(client, ha, {"key_issue": mid})
    done = client.post(f"/reports/{rid}/finalize", json={"exception_reason": "내부 점검용 예외"}, headers=hb)
    assert done.status_code == 200
    doc, footers = _docx_of(client, ha, rid)
    texts = [p.text for p in doc.paragraphs]
    assert any("원고 미배치" in t for t in texts)  # 필수 파트가 비어 있으면 예외 승인 사실과 함께 그대로 보인다
    assert "출력용 확정본(예외 승인)" in footers


# ---------------------------------------------------------------- 간결형 요약본 ----
def test_compact_summary_has_no_cover_or_toc_and_is_only_for_summary(client, db_session):
    make_user(db_session, "c1@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "c1@example.com")
    assert client.post("/reports", json={"report_type": "general", "title": "x", "layout": "compact"}, headers=h).status_code == 422
    mid = _make_manuscript(client, h)
    created = client.post("/reports", json={"report_type": "summary", "title": "간결형 요약본", "layout": "compact"}, headers=h).json()
    assert created["layout"] == "compact"
    rid = created["id"]
    sentence = next(s for s in client.get(f"/reports/source-sentences/{mid}", headers=h).json()["sentences"] if "제정안" in s)
    assert client.post(f"/reports/{rid}/summary-items", headers=h,
                       json={"part_key": "key_issue", "manuscript_id": mid, "source_sentence": sentence}).status_code == 201
    doc, footers = _docx_of(client, h, rid)
    texts = [p.text for p in doc.paragraphs]
    assert "목  차" not in texts  # 표지·독립 목차 없음
    assert doc.paragraphs[0].style.name == "Title" and "간결형 요약본" in doc.paragraphs[0].text
    assert any("임시 간결형 요약본" in t for t in texts) and "임시" in "".join(texts)
    assert any(sentence[:10] in t for t in texts)  # 선택한 핵심 문장은 그대로 들어간다
    assert "내부 검토용 초안" in footers  # 첫 쪽에도 상태 표시가 보인다


@requires_soffice
@pytest.mark.render
def test_compact_summary_is_short_when_rendered(client, db_session):
    make_user(db_session, "c2@example.com", UserRole.RESEARCHER)
    h = auth_headers(client, "c2@example.com")
    mid = _make_manuscript(client, h)
    rid = client.post("/reports", json={"report_type": "summary", "title": "간결형", "layout": "compact"}, headers=h).json()["id"]
    sentence = next(s for s in client.get(f"/reports/source-sentences/{mid}", headers=h).json()["sentences"] if "제정안" in s)
    client.post(f"/reports/{rid}/summary-items", headers=h, json={"part_key": "key_issue", "manuscript_id": mid, "source_sentence": sentence})
    pdf = client.post(f"/reports/{rid}/exports", json={"file_format": "pdf"}, headers=h).json()
    assert pdf["status"] == "success", pdf["error_message"]
    assert pdf["page_count"] <= 2  # 정식형은 4쪽이던 요약본이 1~2쪽으로 줄어든다
